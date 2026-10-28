"""Personal inspection access with explicit, restricted serializers at the tenant boundary."""

import json
import secrets
from contextlib import contextmanager
from datetime import date, timedelta
from io import BytesIO
from typing import Annotated, Literal
from uuid import UUID, uuid4
from zoneinfo import ZoneInfo

from fastapi import APIRouter, HTTPException, Query, Request, Response
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import FileResponse
from openpyxl import Workbook
from pydantic import Field
from sqlalchemy import Date, cast, delete, func, select, text, update
from sqlalchemy.orm import Session

from app.appeal_events import CommentInput, add_comment, event_page, record_event, unanswered_comments
from app.appeal_search import appeal_search
from app.auth import AdminUser, Db, digest, hasher
from app.config import base_settings
from app.day_verification import confirmed_condition
from app.db import database_engine
from app.domain import audit, now
from app.inspection_auth import InspectorUser, inspector_dict, lock_inspector
from app.media import media_path
from app.models import (
    AppealPhoto,
    CitizenAppeal,
    District,
    InspectionAgency,
    InspectionGrant,
    InspectionQR,
    InspectionSession,
    Inspector,
    Organization,
    OrganizationLogin,
    Photo,
    User,
)
from app.organization_context import current_organization, organization_scope
from app.organization_registry import control_session, scope_for
from app.organizations import Owner
from app.photo_ordering import photo_order, photo_position, sort_keys
from app.reports import literal
from app.schemas import StrictModel
from app.work_types import work_catalog

router = APIRouter(prefix="/api/inspection", tags=["inspection"])
owner_router = APIRouter(prefix="/api/platform/inspection", tags=["inspection-owner"])
admin_router = APIRouter(prefix="/api/inspection-admin", tags=["inspection-admin"])


class AgencyInput(StrictModel):
    name: str = Field(min_length=1, max_length=150)
    active: bool = True


class InspectorAppealStatusInput(StrictModel):
    status: Literal["COMPLETED", "CANCELLED"]
    expected_version: int = Field(ge=1)


class AccountInput(StrictModel):
    name: str = Field(min_length=1, max_length=150)
    login: str = Field(min_length=3, max_length=100, pattern=r"^[A-Za-z0-9_.@-]+$")
    password: str = Field(min_length=12, max_length=256)


class AccountUpdate(StrictModel):
    name: str | None = Field(default=None, min_length=1, max_length=150)
    active: bool | None = None
    password: str | None = Field(default=None, min_length=12, max_length=256)


def account_view(db, account):
    return inspector_dict(account, db.get(InspectionAgency, account.agency_id))


def add_account(db, body, agency, *, creator=None):
    login = body.login.lower()
    if db.scalar(select(Inspector.id).where(Inspector.login == login)) or db.scalar(
        select(User.id).where(User.login == login)
    ) or db.scalar(select(OrganizationLogin.user_id).where(OrganizationLogin.login == login).limit(1)):
        raise HTTPException(409, "Этот логин уже занят")
    account = Inspector(agency_id=agency, login=login, name=body.name, password_hash=hasher.hash(body.password),
                        active=True, can_create_accounts=creator is None, created_by=creator)
    db.add(account)
    db.flush()
    return account


@owner_router.get("/agencies")
def agencies(db: Db, user: Owner):
    return [{"id": a.id, "name": a.name, "active": a.active,
             "accounts": db.scalar(select(func.count()).select_from(Inspector).where(Inspector.agency_id == a.id))}
            for a in db.scalars(select(InspectionAgency).order_by(InspectionAgency.created_at, InspectionAgency.id))]


@owner_router.post("/agencies", status_code=201)
def create_agency(body: AgencyInput, db: Db, user: Owner):
    agency = InspectionAgency(id=str(uuid4()), **body.model_dump())
    db.add(agency)
    audit(db, user.id, "INSPECTION_AGENCY_CREATED", "inspection_agency", 0, None, {"id": agency.id, "name": agency.name})
    db.commit()
    return {"id": agency.id, "name": agency.name, "active": agency.active}


@owner_router.patch("/agencies/{agency_id}")
def update_agency(agency_id: UUID, body: AgencyInput, db: Db, user: Owner):
    agency = db.get(InspectionAgency, str(agency_id), with_for_update=True)
    if not agency:
        raise HTTPException(404, "Проверяющая администрация не найдена")
    agency.name, agency.active = body.name, body.active
    if not agency.active:
        db.execute(delete(InspectionSession).where(InspectionSession.inspector_id.in_(
            select(Inspector.id).where(Inspector.agency_id == agency.id))))
    audit(db, user.id, "INSPECTION_AGENCY_UPDATED", "inspection_agency", 0, None, body.model_dump() | {"id": agency.id})
    db.commit()
    return {"ok": True}


@owner_router.get("/agencies/{agency_id}/accounts")
def owner_accounts(agency_id: UUID, db: Db, user: Owner):
    return [account_view(db, a) for a in db.scalars(select(Inspector).where(
        Inspector.agency_id == str(agency_id)).order_by(Inspector.id))]


@owner_router.post("/agencies/{agency_id}/accounts", status_code=201)
def owner_create_account(agency_id: UUID, body: AccountInput, db: Db, user: Owner):
    agency = db.get(InspectionAgency, str(agency_id))
    if not agency:
        raise HTTPException(404, "Проверяющая администрация не найдена")
    account = add_account(db, body, agency.id)
    audit(db, user.id, "INSPECTOR_CREATED", "inspector", account.id, None, {"agency_id": agency.id, "login": account.login})
    db.commit()
    return account_view(db, account)


@owner_router.patch("/accounts/{account_id}")
def owner_update_account(account_id: int, body: AccountUpdate, db: Db, user: Owner):
    lock_inspector(db, account_id)
    account = db.get(Inspector, account_id, with_for_update=True)
    if not account:
        raise HTTPException(404, "Проверяющий не найден")
    for key in ("name", "active"):
        if getattr(body, key) is not None:
            setattr(account, key, getattr(body, key))
    if body.password:
        account.password_hash = hasher.hash(body.password)
    if body.password or body.active is False:
        db.execute(delete(InspectionSession).where(InspectionSession.inspector_id == account.id))
    audit(db, user.id, "INSPECTOR_UPDATED", "inspector", account.id, None,
          {"name": account.name, "active": account.active, "password_reset": bool(body.password)})
    db.commit()
    return account_view(db, account)


@owner_router.get("/accounts/{account_id}/grants")
def owner_grants(account_id: int, db: Db, user: Owner):
    return [{"id": org.id, "name": org.name, "created_at": grant.created_at} for grant, org in db.execute(
        select(InspectionGrant, Organization).join(Organization).where(InspectionGrant.inspector_id == account_id))]


@owner_router.delete("/accounts/{account_id}/grants/{organization_id}")
def owner_revoke(account_id: int, organization_id: UUID, db: Db, user: Owner):
    db.execute(delete(InspectionGrant).where(InspectionGrant.inspector_id == account_id,
                                            InspectionGrant.organization_id == str(organization_id)))
    audit(db, user.id, "INSPECTION_ACCESS_REVOKED", "inspector", account_id, None,
          {"organization_id": str(organization_id)})
    db.commit()
    return {"ok": True}


@router.get("/me")
def me(user: InspectorUser):
    return inspector_dict(user, user._agency)


@router.post("/logout")
def logout(request: Request, response: Response, db: Db, user: InspectorUser):
    db.delete(request.state.inspection_session)
    db.commit()
    for name in ("inspection_session", "inspection_csrf"):
        response.delete_cookie(name, path="/")
    return {"ok": True}


@router.get("/accounts")
def accounts(db: Db, user: InspectorUser):
    if not user.can_create_accounts:
        raise HTTPException(403, "Создавать аккаунты могут только проверяющие, добавленные владельцем")
    return [account_view(db, a) for a in db.scalars(select(Inspector).where(
        Inspector.created_by == user.id).order_by(Inspector.id))]


@router.post("/accounts", status_code=201)
def create_account(body: AccountInput, db: Db, user: InspectorUser):
    if not user.can_create_accounts:
        raise HTTPException(403, "Создавать аккаунты могут только проверяющие, добавленные владельцем")
    account = add_account(db, body, user.agency_id, creator=user.id)
    audit(db, None, "INSPECTOR_CREATED_BY_INSPECTOR", "inspector", account.id, None, {"created_by": user.id})
    db.commit()
    return account_view(db, account)


def issue_qr(db, organization_id):
    org = db.get(Organization, organization_id, with_for_update=True)
    if not org or org.status != "ACTIVE":
        raise HTTPException(409, "Организация недоступна")
    db.execute(update(InspectionQR).where(InspectionQR.organization_id == org.id,
                                         InspectionQR.used_by.is_(None)).values(revoked=True))
    token = secrets.token_urlsafe(32)
    expires = now() + timedelta(minutes=15)
    db.add(InspectionQR(token_hash=digest(token), organization_id=org.id, expires_at=expires, revoked=False))
    # The secret is in the fragment: it never enters access logs or Referer headers.
    url = base_settings().app_origin.rstrip("/") + "/inspection#invite=" + token
    return {"url": url, "expires_at": expires}


@owner_router.post("/organizations/{organization_id}/qr")
def owner_qr(organization_id: UUID, db: Db, user: Owner):
    result = issue_qr(db, str(organization_id))
    audit(db, user.id, "INSPECTION_QR_CREATED", "organization", 0, None, {"id": str(organization_id)})
    db.commit()
    return result


@admin_router.post("/qr")
def admin_qr(user: AdminUser):
    org = current_organization()
    if not org:
        raise HTTPException(404)
    with control_session() as db:
        result = issue_qr(db, org.id)
        audit(db, None, "INSPECTION_QR_CREATED", "organization", 0, None, {"id": org.id, "admin_id": user.id})
        db.commit()
        return result


class QRInput(StrictModel):
    token: str = Field(min_length=40, max_length=100, pattern=r"^[A-Za-z0-9_-]+$")


@router.post("/redeem")
def redeem(body: QRInput, db: Db, user: InspectorUser):
    invitation = db.get(InspectionQR, digest(body.token), with_for_update=True)
    if not invitation or invitation.revoked or invitation.used_by is not None or invitation.expires_at <= now():
        raise HTTPException(409, "QR уже использован, заменён или истёк. Попросите новый QR организации.")
    org = db.get(Organization, invitation.organization_id)
    if not org or org.status != "ACTIVE":
        raise HTTPException(404, "Организация недоступна")
    # Lock the account as well: distinct invitations redeemed simultaneously cannot duplicate a grant.
    db.get(Inspector, user.id, with_for_update=True)
    if not db.get(InspectionGrant, (user.id, org.id)):
        db.add(InspectionGrant(inspector_id=user.id, organization_id=org.id))
    invitation.used_by, invitation.used_at = user.id, now()
    audit(db, None, "INSPECTION_ACCESS_GRANTED", "inspector", user.id, None, {"organization_id": org.id})
    db.commit()
    return {"id": org.id, "name": org.name}


@router.get("/organizations")
def organizations(db: Db, user: InspectorUser):
    return [{"id": org.id, "name": org.name, "city": org.city} for org in db.scalars(select(Organization).join(
        InspectionGrant, InspectionGrant.organization_id == Organization.id).where(
            InspectionGrant.inspector_id == user.id, Organization.status == "ACTIVE").order_by(Organization.name))]


def authorized_organization(db, user, organization_id):
    org = db.get(Organization, str(organization_id))
    if not org or org.status != "ACTIVE" or not db.get(InspectionGrant, (user.id, org.id)):
        raise HTTPException(404, "Организация недоступна. Отсканируйте её новый QR.")
    return org


@contextmanager
def tenant_reader(control, user, organization_id):
    org = authorized_organization(control, user, organization_id)
    scope = scope_for(org)
    with organization_scope(scope), Session(database_engine(scope.database_url)) as tenant:
        # Read-only at the database level as an additional boundary for the inspection service.
        tenant.execute(text("SET TRANSACTION READ ONLY"))
        yield tenant


class InspectionFilters(StrictModel):
    date_from: date | None = None
    date_to: date | None = None
    district_id: int | None = Field(default=None, ge=0)
    work_type: str | None = Field(default=None, max_length=60)
    is_defect: bool | None = None
    photo_id: int | None = Field(default=None, gt=0)
    date_basis: Literal["photo", "received"] = "photo"
    sort_order: Literal["asc", "desc"] = "desc"
    page: int = Field(default=1, ge=1)
    page_size: int = Field(default=40, ge=1, le=100)


def visible_photos(filters):
    stmt = select(Photo).where(Photo.cohort == "production", Photo.preview_removed_at.is_(None), confirmed_condition())
    date_field = Photo.photo_date if filters.date_basis == "photo" else cast(func.timezone(base_settings().app_timezone, Photo.received_at), Date)
    if filters.date_from:
        stmt = stmt.where(date_field >= filters.date_from)
    if filters.date_to:
        stmt = stmt.where(date_field <= filters.date_to)
    if filters.photo_id:
        stmt = stmt.where(Photo.id == filters.photo_id)
    if filters.district_id is not None:
        stmt = stmt.where(Photo.district_id == filters.district_id if filters.district_id else Photo.district_id.is_(None))
    if filters.work_type:
        stmt = stmt.where(Photo.work_type == filters.work_type if filters.work_type != "NONE" else Photo.work_type.is_(None))
    if filters.is_defect is not None:
        stmt = stmt.where(Photo.is_defect == filters.is_defect)
    return stmt


def restricted_photo(photo, districts, catalog, org_id):
    from app.catalog import work_type_label
    # An allowlist is deliberate: future operator fields must never become public here.
    result = {key: getattr(photo, key) for key in (
        "id", "photo_date", "photo_time", "received_at", "latitude", "longitude", "district_id", "work_type", "is_defect",
        "detected_address", "detected_city",
    )}
    result.update(district_name=districts.get(photo.district_id),
                  status="ACCEPTED", original_available=bool(photo.yandex_disk_path),
                  work_type_name=work_type_label(photo.work_type, photo.is_defect, catalog),
                  preview_url=f"/api/inspection/organizations/{org_id}/photos/{photo.id}/preview")
    return result


@router.get("/organizations/{organization_id}/meta")
def metadata(organization_id: UUID, db: Db, user: InspectorUser):
    with tenant_reader(db, user, organization_id) as tenant:
        from app.work_types import catalog_rows
        return {"districts": [{"id": d.id, "name": d.name, "active": d.active} for d in tenant.scalars(select(District).order_by(District.name))],
                "work_types": [{"code": w.code, "name": w.name, "description": w.description} for w in catalog_rows(tenant)],
                "yandex_maps_api_key": base_settings().yandex_maps_api_key,
                "timezone": base_settings().app_timezone, "today": now().astimezone(ZoneInfo(base_settings().app_timezone)).date()}


@router.get("/organizations/{organization_id}/districts")
def district_map(organization_id: UUID, db: Db, user: InspectorUser):
    with tenant_reader(db, user, organization_id) as tenant:
        return {"type": "FeatureCollection", "features": [
            {"type": "Feature", "id": id, "properties": {"id": id, "name": name, "active": active},
             "geometry": json.loads(geometry)} for id, name, active, geometry in tenant.execute(
                 select(District.id, District.name, District.active, func.ST_AsGeoJSON(District.geometry)).order_by(District.name))]}


class MapFilters(InspectionFilters):
    after_id: int = Field(default=0, ge=0)
    snapshot_id: int | None = Field(default=None, ge=0)


@router.get("/organizations/{organization_id}/photos/map")
def photo_map(organization_id: UUID, db: Db, user: InspectorUser, filters: Annotated[MapFilters, Query()]):
    with tenant_reader(db, user, organization_id) as tenant:
        selection = visible_photos(filters)
        snapshot = filters.snapshot_id
        if snapshot is None:
            snapshot = tenant.scalar(selection.with_only_columns(func.max(Photo.id))) or 0
        selection = selection.where(Photo.id <= snapshot)
        located = Photo.latitude.between(-90, 90) & Photo.longitude.between(-180, 180)
        summary = None
        if filters.after_id == 0:
            total, count = tenant.execute(selection.with_only_columns(func.count(), func.count().filter(located))).one()
            summary = {"total": total, "located": count, "without_coordinates": total - count}
        rows = list(tenant.scalars(selection.where(located, Photo.id > filters.after_id).order_by(Photo.id).limit(1001)))
        districts = dict(tenant.execute(select(District.id, District.name)).all())
        catalog = work_catalog(tenant)
        return {"items": [restricted_photo(p, districts, catalog, str(organization_id)) for p in rows[:1000]],
                "summary": summary, "snapshot_id": snapshot,
                "next_after_id": rows[999].id if len(rows) > 1000 else None}


@router.get("/organizations/{organization_id}/photos/arrivals")
def arrivals(organization_id: UUID, db: Db, user: InspectorUser):
    from app.models import DayClosure, DayPhotoReview
    with tenant_reader(db, user, organization_id) as tenant:
        rows = tenant.scalars(visible_photos(InspectionFilters()).join(DayPhotoReview, DayPhotoReview.photo_id == Photo.id)
            .join(DayClosure, DayClosure.id == DayPhotoReview.closure_id)
            .where(DayClosure.confirmed_at > now() - timedelta(days=1)).with_only_columns(Photo.id).order_by(Photo.id))
        return {"photo_ids": list(rows)}


@router.get("/organizations/{organization_id}/photos")
def photos(organization_id: UUID, db: Db, user: InspectorUser, filters: Annotated[InspectionFilters, Query()]):
    with tenant_reader(db, user, organization_id) as tenant:
        stmt = visible_photos(filters)
        total = tenant.scalar(select(func.count()).select_from(stmt.subquery()))
        page = min(filters.page, max(1, (total + filters.page_size - 1) // filters.page_size))
        rows = tenant.scalars(stmt.order_by(*photo_order(filters))
                              .offset((page - 1) * filters.page_size).limit(filters.page_size))
        districts = dict(tenant.execute(select(District.id, District.name)).all())
        catalog = work_catalog(tenant)
        return {"items": [restricted_photo(p, districts, catalog, str(organization_id)) for p in rows],
                "page": page, "page_size": filters.page_size, "total": total}


def require_visible_photo(tenant, photo_id):
    photo = tenant.scalar(visible_photos(InspectionFilters()).where(Photo.id == photo_id))
    if photo is None:
        raise HTTPException(404, "Фотография недоступна до закрытия дня или после изменения данных")
    return photo


@router.get("/organizations/{organization_id}/photos/{photo_id}")
def photo_detail(organization_id: UUID, photo_id: int, db: Db, user: InspectorUser):
    with tenant_reader(db, user, organization_id) as tenant:
        photo = require_visible_photo(tenant, photo_id)
        return restricted_photo(photo, dict(tenant.execute(select(District.id, District.name)).all()),
                                work_catalog(tenant), str(organization_id))


@router.get("/organizations/{organization_id}/photos/{photo_id}/preview")
def photo_preview(organization_id: UUID, photo_id: int, db: Db, user: InspectorUser):
    with tenant_reader(db, user, organization_id) as tenant:
        require_visible_photo(tenant, photo_id)
        path = media_path(photo_id, True)
        if not path.is_file():
            raise HTTPException(404, "Превью временно недоступно")
        return FileResponse(path, media_type="image/jpeg", headers={"Cache-Control": "private, no-store"})


@router.get("/organizations/{organization_id}/photos/{photo_id}/navigation")
def photo_navigation(organization_id: UUID, photo_id: int, db: Db, user: InspectorUser,
                     filters: Annotated[InspectionFilters, Query()]):
    with tenant_reader(db, user, organization_id) as tenant:
        anchor = require_visible_photo(tenant, photo_id)
        selection = visible_photos(filters)
        previous_order = [getattr(col, "desc" if direction == "asc" else "asc")().nullsfirst()
                          for col, direction in sort_keys(filters)]
        return {"next": tenant.scalar(selection.where(photo_position(filters, anchor)).with_only_columns(Photo.id)
                                       .order_by(*photo_order(filters)).limit(1)),
                "previous": tenant.scalar(selection.where(photo_position(filters, anchor, before=True)).with_only_columns(Photo.id)
                                           .order_by(*previous_order).limit(1))}


@router.get("/organizations/{organization_id}/photos/{photo_id}/original")
def original(organization_id: UUID, photo_id: int, db: Db, user: InspectorUser):
    import httpx

    from app.integrations import YandexDisk, safe_url
    with tenant_reader(db, user, organization_id) as tenant:
        photo = require_visible_photo(tenant, photo_id)
        extension = photo.extension if photo.extension in {".jpg", ".jpeg", ".png", ".webp", ".heic", ".tiff", ".gif", ".bmp"} else ".bin"
        headers = {"Cache-Control": "private, no-store", "Content-Disposition": f'inline; filename="photo-{photo.id}{extension}"'}
        local = media_path(photo.id)
        if local.is_file():
            return FileResponse(local, media_type=photo.mime_type or "application/octet-stream", headers=headers)
        if not photo.yandex_disk_path:
            raise HTTPException(404, "Оригинал пока недоступен")
        try:
            result = YandexDisk().request("GET", "/resources/download", params={"path": photo.yandex_disk_path})
            result.raise_for_status()
            url = safe_url(result.json()["href"], ("yandex.net", "yandex.ru", "yandex.com"))
            content = bytearray()
            with httpx.stream("GET", url, timeout=60, follow_redirects=False) as response:
                response.raise_for_status()
                for chunk in response.iter_bytes(64 * 1024):
                    if len(content) + len(chunk) > base_settings().max_file_bytes:
                        raise ValueError("Original too large")
                    content.extend(chunk)
        except (httpx.HTTPError, ValueError, KeyError, RuntimeError, OSError) as exc:
            raise HTTPException(502, "Оригинал временно недоступен. Повторите позже.") from exc
        return Response(bytes(content), media_type=photo.mime_type or "application/octet-stream", headers=headers)


def inspection_report(tenant, filters):
    source = visible_photos(filters).subquery()
    from app.catalog import work_type_label
    catalog = work_catalog(tenant)
    districts = [{"id": did or 0, "name": name or "Не определён", "count": count} for did, name, count in tenant.execute(
        select(source.c.district_id, District.name, func.count()).outerjoin(District, District.id == source.c.district_id)
        .group_by(source.c.district_id, District.name).order_by(func.count().desc()))]
    works = [{"id": code or "NONE", "name": work_type_label(code, catalog=catalog) or "Не определён", "count": count}
             for code, count in tenant.execute(select(source.c.work_type, func.count()).group_by(source.c.work_type)
                                              .order_by(func.count().desc()))]
    matrix = [{"work_type": code or "NONE", "district_id": district or 0, "count": count} for code, district, count in tenant.execute(
        select(source.c.work_type, source.c.district_id, func.count()).group_by(source.c.work_type, source.c.district_id))]
    return {"total": tenant.scalar(select(func.count()).select_from(source)), "districts": districts, "work_district_matrix": matrix,
            "work_types": works, "defects": tenant.scalar(select(func.count()).select_from(source).where(source.c.is_defect))}


@router.get("/organizations/{organization_id}/reports")
def reports(organization_id: UUID, db: Db, user: InspectorUser, filters: Annotated[InspectionFilters, Query()]):
    with tenant_reader(db, user, organization_id) as tenant:
        return inspection_report(tenant, filters)


@router.get("/organizations/{organization_id}/export.xlsx")
def export(organization_id: UUID, db: Db, user: InspectorUser, filters: Annotated[InspectionFilters, Query()]):
    with tenant_reader(db, user, organization_id) as tenant:
        districts = dict(tenant.execute(select(District.id, District.name)).all())
        catalog = work_catalog(tenant)
        wb = Workbook(write_only=True)
        sheet = wb.create_sheet("Фотографии")
        sheet.append(["ID", "Дата", "Время", "Микрорайон", "Вид работы", "Широта", "Долгота", "Адрес", "Город", "Дефект"])
        for photo in tenant.scalars(visible_photos(filters).order_by(Photo.photo_date, Photo.photo_time, Photo.id)
                                    .execution_options(yield_per=500)):
            p = restricted_photo(photo, districts, catalog, str(organization_id))
            sheet.append([literal(p[k]) for k in ("id", "photo_date", "photo_time", "district_name", "work_type_name",
                                                  "latitude", "longitude", "detected_address", "detected_city", "is_defect")])
        report = inspection_report(tenant, filters)
        for key, title in (("districts", "По микрорайонам"), ("work_types", "По видам работ")):
            sheet = wb.create_sheet(title)
            sheet.append(["Название", "Количество"])
            for row in report[key]:
                sheet.append([literal(row["name"]), row["count"]])
        sheet = wb.create_sheet("Работы и микрорайоны")
        sheet.append(["Вид работы"] + [literal(d["name"]) for d in report["districts"]] + ["Всего"])
        matrix = {(r["work_type"], r["district_id"]): r["count"] for r in report["work_district_matrix"]}
        for work in report["work_types"]:
            sheet.append([literal(work["name"])] + [matrix.get((work["id"], district["id"]), 0)
                                                    for district in report["districts"]] + [work["count"]])
        buf = BytesIO()
        wb.save(buf)
        return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Cache-Control": "private, no-store", "Content-Disposition": 'attachment; filename="inspection.xlsx"'})


@router.get("/appeals/notifications")
def all_appeal_notifications(db: Db, user: InspectorUser):
    allowed = select(InspectionGrant.organization_id).join(Organization, Organization.id == InspectionGrant.organization_id).where(
        InspectionGrant.inspector_id == user.id, Organization.status == "ACTIVE")
    condition = CitizenAppeal.organization_id.in_(allowed) & CitizenAppeal.archived_at.is_(None)
    latest = db.scalar(select(CitizenAppeal).where(condition, CitizenAppeal.status == "NEW")
                       .order_by(CitizenAppeal.id.desc()).limit(1))
    return {"new_count": db.scalar(select(func.count()).select_from(CitizenAppeal).where(condition, CitizenAppeal.status == "NEW")),
            "latest_id": db.scalar(select(func.max(CitizenAppeal.id)).where(condition)) or 0,
            "organization_id": latest.organization_id if latest else None}


@router.get("/organizations/{organization_id}/appeals/catalog")
def appeal_catalog(organization_id: UUID, db: Db, user: InspectorUser):
    from app.appeal_catalog import organization_appeal_catalog
    org = authorized_organization(db, user, organization_id)
    return {"organization": {**organization_appeal_catalog(db, org), "color": "#2563eb"},
            "yandex_maps_api_key": base_settings().yandex_maps_api_key}


def inspector_submission(db, user, organization_id, client_id):
    appeal = db.scalar(select(CitizenAppeal).where(CitizenAppeal.submitted_by_inspector_id == user.id,
                                                  CitizenAppeal.inspector_submission_id == str(client_id)))
    if appeal and appeal.organization_id != str(organization_id):
        raise HTTPException(409, "Этот номер отправки уже использован для другой организации")
    return appeal


def inspector_appeal_receipt(appeal):
    return {"id": appeal.id, "number": f"ОБ-{appeal.id:06d}", "is_public": False}


@router.get("/organizations/{organization_id}/appeals/submissions/{client_id}")
def resume_appeal_submission(organization_id: UUID, client_id: UUID, db: Db, user: InspectorUser):
    authorized_organization(db, user, organization_id)
    appeal = inspector_submission(db, user, organization_id, client_id)
    if not appeal:
        raise HTTPException(404, "Отправка ещё не подтверждена")
    return inspector_appeal_receipt(appeal)


def save_inspector_appeal(organization_id, inspector_id, body, photos):
    from app.appeal_workflow import create_appeal
    with control_session() as db:
        # One inspector/client UUID creates one appeal, including simultaneous retries.
        actor = db.get(Inspector, inspector_id, with_for_update=True)
        agency = db.get(InspectionAgency, actor.agency_id) if actor else None
        if not actor or not actor.active or not agency or not agency.active:
            raise HTTPException(403, "Аккаунт проверяющего недоступен")
        if not db.get(InspectionGrant, (actor.id, str(organization_id)), with_for_update=True):
            raise HTTPException(404, "Нет доступа к организации")
        authorized_organization(db, actor, organization_id)
        existing = inspector_submission(db, actor, organization_id, body.client_id)
        if existing:
            return inspector_appeal_receipt(existing)
        if str(body.organization_id) != str(organization_id):
            raise HTTPException(422, "Выберите текущую организацию")
        appeal = create_appeal(db, body, photos, is_public=False, inspector=actor)
        audit(db, None, "INSPECTOR_APPEAL_CREATED", "citizen_appeal", appeal.id, None,
              {"inspector_id": actor.id, "inspector_name": actor.name, "organization_id": appeal.organization_id})
        db.commit()
        return inspector_appeal_receipt(appeal)


@router.post("/organizations/{organization_id}/appeals", status_code=201)
async def create_inspector_appeal(organization_id: UUID, request: Request, db: Db, user: InspectorUser):
    from app.appeal_workflow import InspectorAppealInput, multipart
    authorized_organization(db, user, organization_id)
    body, photos = await multipart(request, InspectorAppealInput)
    return await run_in_threadpool(save_inspector_appeal, organization_id, user.id, body, photos)


@router.get("/organizations/{organization_id}/appeals")
def appeals(organization_id: UUID, db: Db, user: InspectorUser, page: int = Query(default=1, ge=1),
            status: Literal["NEW", "IN_REVIEW", "ASSIGNED", "IN_PROGRESS", "COMPLETED", "REJECTED", "CANCELLED", "RESOLVED", "DECLINED", "SPAM"] | None = None,
            archived: bool = False, q: str = Query(default="", max_length=200)):
    authorized_organization(db, user, organization_id)
    condition = CitizenAppeal.organization_id == str(organization_id)
    condition = condition & (CitizenAppeal.archived_at.is_not(None) if archived else CitizenAppeal.archived_at.is_(None))
    if status:
        condition = condition & (CitizenAppeal.status == status)
    for match in appeal_search(q):
        condition = condition & match
    rows = list(db.scalars(select(CitizenAppeal).where(condition)
                          .order_by(CitizenAppeal.id.desc()).offset((page - 1) * 20).limit(20)))
    messages = unanswered_comments(db, [appeal.id for appeal in rows], user, inspector=True)
    return {"total": db.scalar(select(func.count()).select_from(CitizenAppeal).where(condition)), "page": page, "page_size": 20,
            "items": [{**restricted_appeal(db, a), "unanswered_comments": messages.get(a.id, 0)} for a in rows]}


def restricted_appeal(db, appeal):
    from app.appeal_workflow import workflow_detail
    first = db.scalar(select(AppealPhoto.id).where(AppealPhoto.appeal_id == appeal.id).order_by(AppealPhoto.position).limit(1))
    return {key: getattr(appeal, key) for key in ("id", "organization_id", "organization_name", "city", "address", "place",
            "description", "work_type_name", "category", "status", "version", "created_at", "updated_at", "latitude", "longitude",
            "public_response", "resolved_at", "is_public", "submitted_by_inspector_id", "submitted_by_inspector_name", "archived_at")} | workflow_detail(db, appeal,
                lambda path: f"/api/inspection/organizations/{appeal.organization_id}/appeals" + path, private=False) | {
        "number": f"ОБ-{appeal.id:06d}", "photo_count": appeal.photo_count,
        "can_manage": False, "can_execute": False,
        "can_set_task_status": not appeal.archived_at and appeal.status not in {"RESOLVED", "DECLINED", "SPAM"},
        "preview_url": f"/api/inspection/organizations/{appeal.organization_id}/appeals/{appeal.id}/photos/{first}?variant=preview" if first else None,
        "photos": [{"id": p.id, "url": f"/api/inspection/organizations/{appeal.organization_id}/appeals/{appeal.id}/photos/{p.id}",
                    "preview_url": f"/api/inspection/organizations/{appeal.organization_id}/appeals/{appeal.id}/photos/{p.id}?variant=preview",
                    "original_url": f"/api/inspection/organizations/{appeal.organization_id}/appeals/{appeal.id}/photos/{p.id}?variant=original",
                    **{key: getattr(p, key) for key in ("captured_at", "latitude", "longitude", "accuracy")}}
                   for p in db.scalars(select(AppealPhoto).where(AppealPhoto.appeal_id == appeal.id).order_by(AppealPhoto.position))]}


@router.get("/organizations/{organization_id}/appeals/notifications")
def appeal_notifications(organization_id: UUID, db: Db, user: InspectorUser):
    authorized_organization(db, user, organization_id)
    condition = (CitizenAppeal.organization_id == str(organization_id)) & CitizenAppeal.archived_at.is_(None)
    return {"new_count": db.scalar(select(func.count()).select_from(CitizenAppeal).where(
                condition, CitizenAppeal.status == "NEW")),
            "latest_id": db.scalar(select(func.max(CitizenAppeal.id)).where(condition)) or 0}


@router.get("/organizations/{organization_id}/appeals/{appeal_id}/events")
def appeal_events(organization_id: UUID, appeal_id: int, db: Db, user: InspectorUser,
                  before: int = Query(default=0, ge=0), after: int = Query(default=0, ge=0)):
    authorized_organization(db, user, organization_id)
    appeal = db.get(CitizenAppeal, appeal_id)
    if not appeal or appeal.organization_id != str(organization_id):
        raise HTTPException(404, "Обращение не найдено")
    return event_page(db, appeal, inspector=True, before=before, after=after)


@router.post("/organizations/{organization_id}/appeals/{appeal_id}/comments", status_code=201)
def appeal_comment(organization_id: UUID, appeal_id: int, body: CommentInput, db: Db, user: InspectorUser):
    authorized_organization(db, user, organization_id)
    appeal = db.get(CitizenAppeal, appeal_id, with_for_update=True)
    if not appeal or appeal.organization_id != str(organization_id):
        raise HTTPException(404, "Обращение не найдено")
    return add_comment(db, appeal, body, user, inspector=True)


@router.get("/organizations/{organization_id}/appeals/{appeal_id}/attachments/{photo_id}")
def appeal_attachment(organization_id: UUID, appeal_id: int, photo_id: UUID, db: Db, user: InspectorUser,
                      variant: Literal["content", "preview"] = "content"):
    from app.appeal_workflow import image_response
    from app.models import AppealAttachment
    authorized_organization(db, user, organization_id)
    appeal, photo = db.get(CitizenAppeal, appeal_id), db.get(AppealAttachment, str(photo_id))
    if not appeal or appeal.organization_id != str(organization_id) or not photo or photo.appeal_id != appeal.id:
        raise HTTPException(404)
    return image_response(getattr(photo, variant))


@router.get("/organizations/{organization_id}/appeals/{appeal_id}")
def appeal_detail(organization_id: UUID, appeal_id: int, db: Db, user: InspectorUser):
    authorized_organization(db, user, organization_id)
    appeal = db.get(CitizenAppeal, appeal_id)
    if not appeal or appeal.organization_id != str(organization_id):
        raise HTTPException(404, "Обращение не найдено")
    return {**restricted_appeal(db, appeal),
            "unanswered_comments": unanswered_comments(db, [appeal.id], user, inspector=True).get(appeal.id, 0)}


@router.post("/organizations/{organization_id}/appeals/{appeal_id}/task-status")
def inspector_appeal_status(organization_id: UUID, appeal_id: int, body: InspectorAppealStatusInput, db: Db, user: InspectorUser):
    from app.appeal_workflow import latest_task, versioned
    authorized_organization(db, user, organization_id)
    appeal = db.get(CitizenAppeal, appeal_id, with_for_update=True)
    if not appeal or appeal.organization_id != str(organization_id):
        raise HTTPException(404, "Обращение не найдено")
    versioned(appeal, body.expected_version)
    if appeal.archived_at or appeal.status in {"RESOLVED", "DECLINED", "SPAM"}:
        raise HTTPException(409, "Статус закрытого обращения или спама изменить нельзя")
    if appeal.status != body.status:
        before = appeal.status
        task = latest_task(db, appeal.id)
        if task:
            task.status = body.status
        appeal.status, appeal.updated_at = body.status, now()
        appeal.version += 1
        record_event(db, appeal, "STATUS", user, inspector=True)
        audit(db, None, "INSPECTOR_APPEAL_STATUS", "citizen_appeal", appeal.id, {"status": before},
              {"status": appeal.status, "inspector_id": user.id, "organization_id": appeal.organization_id})
    db.commit()
    return {"ok": True, "version": appeal.version}


@router.get("/organizations/{organization_id}/appeals/{appeal_id}/photos/{photo_id}")
def appeal_photo(organization_id: UUID, appeal_id: int, photo_id: UUID, db: Db, user: InspectorUser,
                 variant: Literal["stamped", "preview", "original"] = "stamped"):
    authorized_organization(db, user, organization_id)
    appeal = db.get(CitizenAppeal, appeal_id)
    photo = db.get(AppealPhoto, str(photo_id))
    if not appeal or appeal.organization_id != str(organization_id) or not photo or photo.appeal_id != appeal.id:
        raise HTTPException(404, "Фотография не найдена")
    return Response(getattr(photo, variant), media_type="image/jpeg", headers={"Cache-Control": "private, no-store"})
