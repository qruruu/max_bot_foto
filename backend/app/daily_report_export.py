"""Familiar customer letters with an unbounded row count and an explicit draft mark."""

import math
import re
import textwrap
from datetime import timedelta
from html import escape
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from app.daily_report_schemas import ReportingConfig, ReportingDocument
from app.reports import literal


def letter_data(report, draft=False):
    config = ReportingConfig.model_validate(report.config if draft else report.approved_config)
    document = ReportingDocument.model_validate(report.draft if draft else report.approved)
    works = {w.id: w for w in config.works}
    areas = {a.id: a for a in config.areas}
    letters = []
    for contract in config.contracts:
        blocks = []
        for kind in ("fact", "plan"):
            rows = []
            for line in getattr(document, kind):
                if line.decision not in {"include", "info"} or areas[line.area_id].contract_id != contract.id:
                    continue
                work, area = works[line.work_id], areas[line.area_id]
                rows.append({"work": work.name, "area": area.name, "resources": line.resources,
                             "quantity": line.quantity if line.decision == "include" else None,
                             "unit": work.unit, "note": line.note or ("Без указания объёма" if line.decision == "info" else "")})
            blocks.append(rows)
        letters.append({"contract": contract, "number": report.letter_numbers.get(contract.id, ""),
                        "fact": blocks[0], "plan": blocks[1]})
    return config, document, letters


def export_xlsx(report, *, draft=False):
    config, document, letters = letter_data(report, draft)
    wb = Workbook()
    wb.remove(wb.active)
    profile = config.profile
    summary = wb.create_sheet("Настройки")
    values = [("Дата письма", report.day), ("Факт за день", report.day - timedelta(days=1)),
              ("План на день", report.day), ("Организация", profile.company_name),
              ("Режим", "Тестовый" if report.cohort == "test" else "Рабочий"),
              ("Статус", "ЧЕРНОВИК" if draft else "Утверждён"), ("Версия", report.version if draft else report.approved_version),
              ("Фотоотчёт", profile.photo_url), ("Примечание", document.note)]
    for label, value in values:
        summary.append([literal(label), literal(value)])
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 85
    for row in summary:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if not isinstance(cell.value, str) and hasattr(cell.value, "year"):
                cell.number_format = "dd.mm.yyyy"
    for index, letter in enumerate(letters, 1):
        contract = letter["contract"]
        suffix = re.sub(r"[\\/*?:\[\]]", "-", contract.number)[-20:]
        sheet = wb.create_sheet((f"Письмо {index} " + suffix)[:31])
        sheet.sheet_view.showGridLines = False
        for column, width in zip("ABCDEF", [5.5, 22, 27, 14, 11, 16], strict=True):
            sheet.column_dimensions[column].width = width
        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins = PageMargins(left=.3, right=.3, top=.35, bottom=.35, header=.15, footer=.15)
        row = 1
        def wide(value, *, bold=False, center=False, height=24, size=11, sheet=sheet):
            nonlocal row
            for part in textwrap.wrap(str(value), width=1600, replace_whitespace=False, drop_whitespace=False) or [""]:
                sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
                cell = sheet.cell(row, 1, literal(part))
                cell.font = Font(name="Times New Roman", size=size, bold=bold)
                cell.alignment = Alignment(wrap_text=True, horizontal="center" if center else "left", vertical="center")
                lines = max(1, math.ceil(len(part) / (84 * 11 / size)))
                sheet.row_dimensions[row].height = max(height, lines * size * 1.4)
                row += 1
        if draft or report.cohort == "test":
            wide("ЧЕРНОВИК" if draft else "ТЕСТОВЫЙ ОТЧЁТ", bold=True, center=True)
        wide(profile.company_name, bold=True, center=True, height=36, size=13)
        wide(profile.address, height=28, size=10)
        wide("   ".join(v for v in [f"ИНН {profile.inn}" if profile.inn else "", f"ОГРНИП/ОГРН {profile.ogrn}" if profile.ogrn else "", profile.phone, profile.email] if v), size=9)
        wide(f"№ {letter['number'] or 'без номера (черновик)'} от {report.day:%d.%m.%Y}")
        for recipient in profile.recipient:
            wide(recipient)
        wide(f"В рамках исполнения контракта {contract.number} сообщаю следующую информацию:", height=34)
        border = Border(*(Side(style="thin", color="444444") for _ in range(4)))
        for kind, title, day in [("fact", "Фактическая дислокация оказания услуг (производства работ)", report.day - timedelta(days=1)),
                                  ("plan", "Плановая дислокация оказания услуг", report.day)]:
            wide(title, bold=True, center=True, height=32)
            wide(profile.subject, center=True, height=30)
            wide(f"{profile.short_name or profile.company_name} · {'За' if kind == 'fact' else 'На'} {day:%d.%m.%Y}", center=True)
            headers = ["№", "Наименование работ, услуг", "Место оказания услуг (производства работ)", "Спецтехника, персонал", "Объём" if kind == "fact" else "План, объём", "Примечания"]
            for col, value in enumerate(headers, 1):
                cell = sheet.cell(row, col, value)
                cell.font = Font(name="Times New Roman", size=10, bold=True)
                cell.fill = PatternFill("solid", fgColor="E8ECF0")
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
                cell.border = border
            sheet.row_dimensions[row].height = 42
            row += 1
            lines = letter[kind]
            detailed_notes = []
            if not lines:
                wide("Работы не заявлены" + (": " + document.note if document.note else ""), height=34)
            for n, line in enumerate(lines, 1):
                shown = dict(line)
                for field, limit, label in [("area", 600, "Место оказания услуг"), ("resources", 250, "Персонал и техника"), ("note", 300, "Примечание")]:
                    if len(shown[field]) > limit:
                        detailed_notes.append(f"Строка {n}. {label}: {shown[field]}")
                        shown[field] = f"См. подробное примечание к строке {n} под таблицей"
                for col, value in enumerate([n, shown["work"], shown["area"], shown["resources"], shown["quantity"], shown["note"]], 1):
                    cell = sheet.cell(row, col, literal(value))
                    cell.font = Font(name="Times New Roman", size=10)
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                    cell.border = border
                    if col == 5 and value is not None:
                        cell.number_format = ('0' if float(value).is_integer() else '0.###') + ' "' + line["unit"].replace('"', '""') + '"'
                # Expanded rows are retained and may span multiple printed pages.
                wraps = max(len(shown["work"]) / 26, len(shown["area"]) / 31,
                            len(shown["resources"]) / 17, len(shown["note"]) / 18, 2)
                sheet.row_dimensions[row].height = min(400, math.ceil(wraps) * 14)
                row += 1
            for note in detailed_notes:
                wide(note, size=10)
            row += 1
        if profile.photo_url:
            wide("Фотоотчёт доступен по ссылке " + profile.photo_url, height=42, size=10)
        if document.note:
            wide("Примечание: " + document.note, size=10)
        wide(profile.short_name or profile.company_name, height=40)
        sheet.print_area = f"A1:F{row}"
        sheet.oddFooter.center.text = "Страница &P из &N"
    result = BytesIO()
    wb.save(result)
    return result.getvalue()


def export_html(report, *, draft=False):
    config, document, letters = letter_data(report, draft)
    profile = config.profile
    def esc(value):
        return escape(str(value if value is not None else ""), quote=True)
    content = []
    for letter in letters:
        parts = ["<article>"]
        if draft or report.cohort == "test":
            parts.append("<div class='draft'>" + ("ЧЕРНОВИК" if draft else "ТЕСТОВЫЙ ОТЧЁТ") + "</div>")
        parts += [f"<header><h1>{esc(profile.company_name)}</h1><p>{esc(profile.address)}</p>",
                  f"<p>ИНН {esc(profile.inn)} · ОГРНИП/ОГРН {esc(profile.ogrn)} · {esc(profile.phone)} · {esc(profile.email)}</p></header>",
                  f"<div class='recipient'><p>№ {esc(letter['number'] or 'без номера (черновик)')} от {report.day:%d.%m.%Y}</p><div>" + "<br>".join(map(esc, profile.recipient)) + "</div></div>",
                  f"<p>В рамках исполнения контракта {esc(letter['contract'].number)} сообщаю следующую информацию:</p>"]
        for kind, title, day in [("fact", "Фактическая дислокация оказания услуг (производства работ)", report.day - timedelta(days=1)), ("plan", "Плановая дислокация оказания услуг", report.day)]:
            parts.append(f"<h2>{title}</h2><div class='center'>{esc(profile.subject)}<br>{esc(profile.short_name)}<br><b>{'За' if kind == 'fact' else 'На'} {day:%d.%m.%Y}</b></div>")
            parts.append("<table><thead><tr><th>№</th><th>Наименование работ, услуг</th><th>Место оказания услуг</th><th>Спецтехника, персонал</th><th>Объём</th><th>Примечания</th></tr></thead><tbody>")
            for n, row in enumerate(letter[kind], 1):
                quantity = f"{row['quantity']:g} {row['unit']}" if row["quantity"] is not None else "—"
                parts.append("<tr>" + "".join("<td>" + esc(value) + "</td>" for value in [n, row["work"], row["area"], row["resources"], quantity, row["note"]]) + "</tr>")
            if not letter[kind]:
                parts.append("<tr><td colspan='6'>Работы не заявлены. " + esc(document.note) + "</td></tr>")
            parts.append("</tbody></table>")
        if profile.photo_url:
            parts.append(f"<p>Фотоотчёт доступен по ссылке <a href='{esc(profile.photo_url)}'>{esc(profile.photo_url)}</a></p>")
        if document.note:
            parts.append("<p>Примечание: " + esc(document.note) + "</p>")
        parts.append("<p class='signature'>" + esc(profile.short_name) + "</p></article>")
        content.append("".join(parts))
    return """<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>Ежедневный отчёт</title><style>
      body{font:12pt 'Times New Roman',serif;background:#eee;margin:0;color:#111}article{background:white;box-sizing:border-box;max-width:210mm;margin:20px auto;padding:15mm}header{text-align:center;border-bottom:2px solid #111}header h1{font-size:14pt}header p{font-size:10pt}h2{font-size:12pt;text-align:center;margin:20px 0 8px}.recipient{display:flex;justify-content:space-between;margin:20px 0}.center{text-align:center;line-height:1.5}table{width:100%;border-collapse:collapse;margin:12px 0 20px;font-size:10pt}td,th{border:1px solid #444;padding:7px 5px;overflow-wrap:anywhere}thead{display:table-header-group}th{background:#e8ecf0}.signature{margin-top:40px}.draft{color:#a33;text-align:center;font-weight:bold}nav{font:14px system-ui;padding:15px;text-align:center}button{padding:10px 20px;cursor:pointer}tr{break-inside:avoid}a{color:inherit;overflow-wrap:anywhere}@page{size:A4;margin:12mm}@media print{body{background:white}nav{display:none}article{max-width:none;margin:0;padding:0;break-after:page}article:last-child{break-after:auto}}
    </style><nav><button onclick="window.print()">Печать / Сохранить в PDF</button><p>Каждый контракт начинается с новой страницы. Подпись и печать проставляются при оформлении.</p></nav>""" + "".join(content) + "</html>"
