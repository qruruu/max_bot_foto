"""One physical grid for the letter, with explicit, indivisible printed rows."""

import base64
from functools import lru_cache
from io import BytesIO

from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break
from PIL import ImageFont

from app.reports import literal

COLUMN_PERCENTAGES = (4, 29, 29, 14, 10, 14)
COLUMN_PIXELS = tuple(700 * part / 100 for part in COLUMN_PERCENTAGES)
PAGE_HEIGHT = 760  # A4 inside 12 mm margins, with a little tolerance for Excel/Calc.
BODY_SIZE = 11


@lru_cache(maxsize=8)
def text_font(size, bold=False):
    candidates = [f"/usr/share/fonts/truetype/liberation2/LiberationSerif-{'Bold' if bold else 'Regular'}.ttf",
                  f"/System/Library/Fonts/Supplemental/Times New Roman{' Bold' if bold else ''}.ttf",
                  f"/usr/share/fonts/truetype/dejavu/DejaVuSerif{'-Bold' if bold else ''}.ttf"]
    for path in candidates:
        try:
            return ImageFont.truetype(path, round(size * 96 / 72))
        except OSError:
            continue
    return ImageFont.load_default(size=round(size * 96 / 72))


def wrap_text(value, pixels, *, size=BODY_SIZE, bold=False):
    """Write actual line breaks so row height does not depend on Excel autofit."""
    font = text_font(size, bold)
    lines = []
    for paragraph in str(value).split("\n"):
        line = ""
        for word in paragraph.split():
            candidate = (line + " " + word).strip()
            if line and font.getlength(candidate) > pixels:
                lines.append(line)
                line = ""
            # Very long URLs/identifiers also fit instead of disappearing beyond a border.
            while font.getlength(word) > pixels:
                end = 1
                while end < len(word) and font.getlength(word[:end + 1]) <= pixels:
                    end += 1
                if line:
                    lines.append(line)
                    line = ""
                lines.append(word[:end])
                word = word[end:]
            line = (line + " " + word).strip()
        lines.append(line)
    return "\n".join(lines)


class LetterSheet:
    def __init__(self, sheet, contract):
        self.sheet, self.contract = sheet, contract
        self.row, self.position = 1, 0
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 90
        for column, pixels in zip("ABCDEF", COLUMN_PIXELS, strict=True):
            sheet.column_dimensions[column].width = (pixels - 5) / 7
        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.page_setup.useFirstPageNumber, sheet.page_setup.firstPageNumber = True, 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins = PageMargins(left=12 / 25.4, right=12 / 25.4, top=12 / 25.4,
                                        bottom=12 / 25.4, header=.15, footer=.15)
        sheet.print_options.horizontalCentered = True
        sheet.oddFooter.center.text = f"Контракт {contract.number} · Страница &P"
        sheet.oddFooter.center.size = 8

    def ensure(self, height):
        if self.position and self.position + height > PAGE_HEIGHT:
            self.sheet.row_breaks.append(Break(id=self.row - 1))
            self.position = 0
            return True
        return False

    def advance(self, height):
        self.sheet.row_dimensions[self.row].height = height
        self.row += 1
        self.position += height

    def space(self, height=6):
        if self.position:
            self.advance(height)

    def text(self, value, *, bold=False, center=False, size=BODY_SIZE, minimum=14, keep=0):
        wrapped = wrap_text(value, sum(COLUMN_PIXELS) - 8, size=size, bold=bold)
        # Excel limits a single row to 409 points. Keep arbitrarily long notes in full.
        lines = wrapped.split("\n")
        for start in range(0, len(lines), 24):
            part = lines[start:start + 24]
            height = max(minimum, len(part) * (size + 2) + 2)
            self.ensure(height + (keep if start + 24 >= len(lines) else 0))
            self.sheet.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=6)
            cell = self.sheet.cell(self.row, 1, literal("\n".join(part)))
            cell.font = Font(name="Times New Roman", size=size, bold=bold)
            cell.alignment = Alignment(wrap_text=True, horizontal="center" if center else "left", vertical="center")
            self.advance(height)

    def image_height(self, value, scale=1, max_height=150):
        picture = Image(BytesIO(base64.b64decode(value.split(",", 1)[1])))
        return min(max_height, picture.height / picture.width * sum(COLUMN_PIXELS) * scale * .75)

    def image(self, value, *, scale=1, max_height=150):
        height = self.image_height(value, scale, max_height)
        self.ensure(height)
        picture = Image(BytesIO(base64.b64decode(value.split(",", 1)[1])))
        natural_height = picture.height / picture.width * sum(COLUMN_PIXELS) * .75
        if scale == 1 and natural_height <= max_height:
            # Both ends belong to the same grid as the tables, not a guessed pixel width.
            picture.anchor = TwoCellAnchor(editAs="oneCell", _from=AnchorMarker(col=0, row=self.row - 1),
                                          to=AnchorMarker(col=6, row=self.row))
        else:
            picture.width = height / .75 * picture.width / picture.height
            picture.height = height / .75
            picture.anchor = f"A{self.row}"
        self.sheet.add_image(picture)
        # A picture occupies ONE row. Page breaks can never bisect the signature or seal.
        self.advance(height)

    def recipient(self, number, day, recipients):
        left = f"№ {number or 'б/н'} от {day:%d.%m.%Y}"
        right = "\n".join(recipients)
        spans = [(1, 3, left), (4, 6, right)]
        wrapped = [(a, b, wrap_text(text, sum(COLUMN_PIXELS[a - 1:b]) - 8, size=BODY_SIZE))
                   for a, b, text in spans]
        height = max(text.count("\n") + 1 for _, _, text in wrapped) * 13 + 2
        self.ensure(height)
        for start, end, value in wrapped:
            self.sheet.merge_cells(start_row=self.row, start_column=start, end_row=self.row, end_column=end)
            cell = self.sheet.cell(self.row, start, literal(value))
            cell.font = Font(name="Times New Roman", size=BODY_SIZE)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        self.advance(height)

    @staticmethod
    def cells(values, *, bold=False):
        displayed = [wrap_text(value, width - 10, size=10.5, bold=bold)
                     for value, width in zip(values, COLUMN_PIXELS, strict=True)]
        height = max(28, max(value.count("\n") + 1 for value in displayed) * 12.5 + 6)
        return displayed, height

    def table_row(self, values, *, header=False, number=None, unit=""):
        displayed, height = self.cells(values, bold=header)
        border = Border(*(Side(style="thin", color="222222") for _ in range(4)))
        for col, value in enumerate(displayed, 1):
            cell = self.sheet.cell(self.row, col, literal(value))
            cell.font = Font(name="Times New Roman", size=10.5, bold=header)
            cell.alignment = Alignment(wrap_text=True, vertical="center",
                                       horizontal="center" if header or col == 1 else "right" if col == 5 else "left")
            cell.border = border
            if header:
                cell.fill = PatternFill("solid", fgColor="EEF0F2")
            elif col == 1:
                cell.value = values[0]
            if col == 5 and number is not None:
                cell.value = number
                numeric = '0' if float(number).is_integer() else '0.###'
                label = f"{number:.3f}".rstrip("0").rstrip(".") if not float(number).is_integer() else str(int(number))
                # Put units on their own line when needed; never replace a number by ####.
                separate = text_font(10.5).getlength(label + " " + unit) > COLUMN_PIXELS[4] - 8
                cell.number_format = numeric + ' "' + ("\n" if separate else "") + unit.replace('"', '""') + '"'
                if text_font(10.5).getlength(label) > COLUMN_PIXELS[4] - 8:
                    cell.alignment = Alignment(vertical="center", horizontal="right", shrink_to_fit=True)
        self.advance(height)

    def finish(self):
        self.sheet.print_area = f"A1:F{self.row - 1}"
