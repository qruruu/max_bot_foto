import hashlib
import runpy
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from alembic.migration import MigrationContext
from alembic.operations import Operations
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import select

from app import ml
from app.catalog import WORK_TYPES
from app.domain import now
from app.media import media_path
from app.models import Audit, Inbox, Job, ModelVersion, Photo, User
from app.processing import sync_photo


def test_existing_labels_models_and_remote_paths_migrate(db, monkeypatch):
    migration = runpy.run_path(
        str(Path(__file__).parents[1] / "alembic/versions/0005_short_work_catalog.py")
    )
    user = User(login="operator", name="Operator", password_hash="unused", role="OPERATOR")
    inbox = Inbox(max_chat_id=123, message_id="migration", payload={})
    db.add_all([user, inbox])
    db.flush()
    codes = list(migration["MERGED_CODES"]) + ["BIN_EMPTYING", None]
    photos, files = [], {}
    for i, code in enumerate(codes):
        old_path = f"disk:/MAX_PHOTOS/2026-04-20/Чат_1/старое_имя_{i}.jpg"
        photo = Photo(
            inbox_id=inbox.id, attachment_index=i, max_chat_id=123, chat_name="Чат 1",
            max_message_id="migration", work_type=code, work_type_source="OPERATOR" if code else None,
            ai_alternatives=["PLAYGROUND_SNOW", "MAF_INSPECTION", "BIN_EMPTYING"],
            ai_result={"raw": {"old_category": code}}, reviewed_by=user.id,
            sha256=hashlib.sha256(str(i).encode()).hexdigest(), photo_date=date(2026, 4, 20),
            analyzed_at=now(), status="SPAM" if code is None else "NEEDS_REVIEW", is_spam=code is None,
            extension=".jpg", yandex_disk_path=old_path, storage_status="SYNCED", stored_filename="old.jpg",
        )
        db.add(photo)
        db.flush()
        Image.new("RGB", (20, 20)).save(media_path(photo.id, True))
        files[old_path] = photo.sha256
        photos.append(photo)
    old_model = ModelVersion(
        status="ACTIVE", dataset_fingerprint="0" * 64, sample_count=100,
        labels_snapshot={"1": {"label": "MAF_INSPECTION"}},
        metrics={"architecture": "resnet18-local-v1"}, supported_classes=["MAF_INSPECTION"],
    )
    db.add(old_model)
    db.commit()
    old_paths = {p.id: p.yandex_disk_path for p in photos}
    with Operations.context(MigrationContext.configure(db.connection())):
        migration["upgrade"]()
    db.commit()
    db.expire_all()

    expected = [migration["MERGED_CODES"].get(code, code) for code in codes]
    assert [p.work_type for p in photos] == expected
    assert [p.ai_result["raw"]["old_category"] for p in photos] == codes
    assert all(p.ai_alternatives == ["PLAYGROUND_MAINTENANCE", "BIN_EMPTYING"] for p in photos)
    assert all(p.yandex_disk_path == old_paths[p.id] and p.storage_status == "PENDING" for p in photos)
    assert all(p.version == 2 for p in photos)
    assert old_model.status == "SUPERSEDED"
    assert old_model.labels_snapshot["1"]["label"] == "MAF_INSPECTION"
    assert ml.latest_evaluated(db) is None
    assert ml.active_model(db) is None
    examples, missing = ml.collect_examples(db)
    assert len(examples) == 7 and missing == 0
    assert {e.label for e in examples} <= set(WORK_TYPES)
    jobs = list(db.scalars(select(Job)))
    assert len(jobs) == len(photos)
    assert all(j.kind == "SYNC" and j.status == "PENDING" for j in jobs)
    audit = db.scalar(select(Audit).where(Audit.photo_id == photos[0].id))
    assert audit.old_value["work_type"] == "MAF_INSPECTION"
    assert audit.new_value["work_type"] == "PLAYGROUND_MAINTENANCE"

    class Disk:
        fail = True

        def move(self, source, target, digest):
            if self.fail:
                raise RuntimeError("offline")
            assert files.get(target, files.get(source)) == digest
            if target not in files:
                files[target] = files.pop(source)

    disk = Disk()
    monkeypatch.setattr("app.processing.YandexDisk", lambda: disk)
    with pytest.raises(RuntimeError, match="offline"):
        sync_photo(photos[0].id)
    db.expire_all()
    assert photos[0].storage_status == "ERROR"
    assert photos[0].yandex_disk_path == old_paths[photos[0].id]
    disk.fail = False
    for photo in photos:
        sync_photo(photo.id)
    db.expire_all()
    for photo in photos:
        assert photo.storage_status == "SYNCED"
        assert "/Чат_1/МР_НЕ_ОПРЕДЕЛЕН/" in photo.yandex_disk_path
        assert photo.yandex_disk_path in files and old_paths[photo.id] not in files
        assert photo.stored_filename == photo.yandex_disk_path.rsplit("/", 1)[-1]
    assert "ПЛОЩАДКИ" in photos[0].stored_filename
    assert photos[-1].status == "SPAM"


def test_api_and_export_use_only_short_labels(admin, db):
    response = admin.get("/api/meta")
    assert response.status_code == 200
    assert [w["name"] for w in response.json()["work_types"]] == [v[0] for v in WORK_TYPES.values()]
    inbox = Inbox(max_chat_id=123, message_id="short-labels", payload={})
    db.add(inbox)
    db.flush()
    photo = Photo(
        inbox_id=inbox.id, attachment_index=0, max_chat_id=123, chat_name="Чат 1",
        max_message_id="short-labels", status="NEEDS_REVIEW", work_type="BIN_EMPTYING",
    )
    db.add(photo)
    db.commit()
    assert admin.get(f"/api/photos/{photo.id}").json()["work_type_name"] == "УРНЫ"
    assert admin.get("/api/reports").json()["work_types"][0]["name"] == "УРНЫ"
    export = admin.get("/api/export.xlsx")
    assert export.status_code == 200
    workbook = load_workbook(BytesIO(export.content))
    assert workbook["Все фото"].cell(2, 7).value == "УРНЫ"
    assert admin.get("/api/photos?work_type=MAF_INSPECTION").status_code == 422
