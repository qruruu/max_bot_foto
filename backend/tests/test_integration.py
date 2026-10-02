from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from io import BytesIO

import pytest
from geoalchemy2.shape import from_shape
from openpyxl import load_workbook
from PIL import Image
from shapely.geometry import MultiPolygon, Polygon
from sqlalchemy import func, select, text

from app.catalog import Status
from app.domain import now
from app.intake import ingest
from app.media import media_path
from app.models import Audit, Chat, District, Inbox, Job, Photo
from app.processing import process_photo, regeo_all, sync_photo
from app.recognition import Classification, ClassificationResponse

pytestmark = pytest.mark.integration


def event(mid="mid1", chat_id=123, attachments=None, forward=False):
    message = {
        "recipient": {"chat_id": chat_id},
        "sender": {"user_id": 8},
        "body": {
            "mid": mid,
            "attachments": attachments
            if attachments is not None
            else [{"type": "image", "payload": {"url": "https://cdn.max.ru/photo"}}],
        },
    }
    if forward:
        message["link"] = {"type": "forward", "message": {"mid": "original"}}
    return {"update_type": "message_created", "message": message}


def setup_chat(db):
    chat = Chat(max_chat_id=123, name="Чат 1", active=True)
    db.add(chat)
    db.commit()
    return chat


def setup_district(db, name="12 мкр", bounds=(72, 61, 73, 62)):
    x1, y1, x2, y2 = bounds
    geom = MultiPolygon([Polygon([(x1, y1), (x2, y1), (x2, y2), (x1, y2), (x1, y1)])])
    district = District(name=name, active=True, geometry=from_shape(geom, srid=4326))
    db.add(district)
    db.commit()
    return district


def intake_photo(db, mid="mid1", **kwargs):
    inbox_id = ingest(db, event(mid=mid, **kwargs))
    db.commit()
    return db.scalar(select(Photo).where(Photo.inbox_id == inbox_id))


class FakeMax:
    def __init__(self, color="green"):
        out = BytesIO()
        Image.new("RGB", (100, 100), color=color).save(out, format="PNG")
        self.content = out.getvalue()
        self.downloads = 0

    def download(self, attachment, target):
        self.downloads += 1
        target.write_bytes(self.content)


class FakeDisk:
    def __init__(self):
        self.files = {}
        self.fail = False
        self.upload_calls = 0

    def upload(self, original, target, digest):
        self.upload_calls += 1
        self.files[target] = original.read_bytes()
        if self.fail:
            raise RuntimeError("Simulated connection loss after remote upload")

    def move(self, source, target, digest):
        if target not in self.files:
            self.files[target] = self.files.pop(source)


class GoodOCR:
    def recognize(self, path):
        return "30.09.2026 11:43:25\n61,1060N 72,5780E\n26 Парковая улица\nНефтеюганск"


class BrokenOCR:
    def recognize(self, path):
        raise RuntimeError("OCR unavailable")


class GoodAI:
    def classify(self, path):
        c = Classification(result="SINGLE", work_type="BIN_EMPTYING", confidence=0.93, alternatives=[])
        return ClassificationResponse(c, {"provider": "test", "result": c.model_dump()})


class MultipleAI:
    def classify(self, path):
        c = Classification(
            result="MULTIPLE", work_type=None, confidence=0.9, alternatives=["BIN_EMPTYING", "SWEEPING"]
        )
        return ClassificationResponse(c, {"result": c.model_dump()})


class BrokenAI:
    def classify(self, path):
        raise RuntimeError("AI unavailable")


def test_actual_postgis_and_unique_hash_schema(db):
    assert "POSTGIS" in db.scalar(text("SELECT PostGIS_Full_Version()"))
    assert (
        db.scalar(
            text("SELECT count(*) FROM pg_indexes WHERE tablename='photos' AND indexdef LIKE '%UNIQUE%'")
        )
        >= 2
    )


def test_webhook_security_allowlist_idempotency(client, db):
    payload = event()
    assert client.post("/api/webhooks/max", json=payload).status_code == 403
    headers = {"X-Max-Bot-Api-Secret": "test-secret-" * 4}
    assert client.post("/api/webhooks/max", json=payload, headers=headers).status_code == 200
    assert db.scalar(select(func.count()).select_from(Inbox)) == 0
    setup_chat(db)
    for _ in range(3):
        assert client.post("/api/webhooks/max", json=payload, headers=headers).status_code == 200
    assert db.scalar(select(func.count()).select_from(Inbox)) == 1
    assert db.scalar(select(func.count()).select_from(Photo)) == 1
    assert db.scalar(select(func.count()).select_from(Job)) == 1


def test_forwarded_and_unsupported_never_download(db):
    setup_chat(db)
    max_client = FakeMax()
    p = intake_photo(db, forward=True)
    process_photo(p.id, max_client=max_client, disk=FakeDisk(), ocr=GoodOCR(), classifier=GoodAI())
    assert max_client.downloads == 0
    db.refresh(p)
    assert p.status == Status.FORWARDED and not p.sha256
    p2 = intake_photo(db, mid="file", attachments=[{"type": "file", "filename": "report.pdf"}])
    assert p2.status == Status.REJECTED
    assert not media_path(p.id).exists()


def test_success_original_preserved_and_fields(db):
    setup_chat(db)
    d = setup_district(db)
    p = intake_photo(db)
    mx = FakeMax()
    disk = FakeDisk()
    process_photo(p.id, mx, disk, GoodOCR(), GoodAI())
    db.refresh(p)
    assert p.status == Status.ACCEPTED
    assert p.district_id == d.id and p.work_type == "BIN_EMPTYING"
    assert p.photo_date == date(2026, 9, 30) and p.photo_time.second == 25
    assert p.detected_city == "Нефтеюганск"
    assert disk.files[p.yandex_disk_path] == mx.content
    assert not media_path(p.id).exists() and media_path(p.id, True).exists()
    assert p.storage_status == "SYNCED"
    assert p.ai_result["raw"]["provider"] == "test"


@pytest.mark.parametrize(
    "ocr,ai,reason",
    [
        (BrokenOCR(), GoodAI(), "DATE_MISSING"),
        (GoodOCR(), BrokenAI(), "WORK_UNCERTAIN"),
        (GoodOCR(), MultipleAI(), "WORK_MULTIPLE"),
    ],
)
def test_recognition_failure_accepts_and_saves(db, ocr, ai, reason):
    setup_chat(db)
    setup_district(db)
    p = intake_photo(db)
    disk = FakeDisk()
    process_photo(p.id, FakeMax(), disk, ocr, ai)
    db.refresh(p)
    assert p.status == Status.NEEDS_REVIEW and reason in p.review_reason
    assert len(disk.files) == 1 and p.yandex_disk_path


def test_outside_overlap_and_boundary_use_only_coordinates(db):
    setup_chat(db)
    first = setup_district(db)
    p = intake_photo(db)
    from app.domain import resolve_district

    assert resolve_district(db, 61, 72) == (first.id, None)  # ST_Covers includes the boundary.
    assert resolve_district(db, 40, 40) == (None, "OUTSIDE_DISTRICTS")
    setup_district(db, "Пересечение", (72.5, 61, 73.5, 62))
    assert resolve_district(db, 61.106, 72.578) == (None, "DISTRICT_OVERLAP")
    process_photo(p.id, FakeMax(), FakeDisk(), GoodOCR(), GoodAI())
    db.refresh(p)
    assert p.status == Status.NEEDS_REVIEW and p.district_id is None


def test_duplicate_globally_and_concurrent(db):
    setup_chat(db)
    setup_district(db)
    db.add(Chat(max_chat_id=456, name="Чат 2", active=True))
    db.commit()
    p1 = intake_photo(db)
    p2 = intake_photo(db, mid="second", chat_id=456)
    disk = FakeDisk()
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(process_photo, p.id, FakeMax(), disk, GoodOCR(), GoodAI()) for p in (p1, p2)]
        for f in futures:
            f.result(timeout=30)
    db.refresh(p1)
    db.refresh(p2)
    assert sorted([p1.status, p2.status]) == ["ACCEPTED", "DUPLICATE"]
    assert len(disk.files) == 1
    duplicate = p1 if p1.status == Status.DUPLICATE else p2
    assert duplicate.sha256 is None and duplicate.duplicate_of
    assert not media_path(duplicate.id).exists()


def test_storage_outage_retains_original_and_analysis(db):
    setup_chat(db)
    setup_district(db)
    p = intake_photo(db)
    disk = FakeDisk()
    disk.fail = True
    with pytest.raises(RuntimeError):
        process_photo(p.id, FakeMax(), disk, GoodOCR(), GoodAI())
    db.refresh(p)
    assert p.status == Status.ERROR and media_path(p.id).exists() and p.analyzed_at
    disk.fail = False
    # Retry must not run recognition again or overwrite its useful results.
    process_photo(p.id, FakeMax(), disk, BrokenOCR(), BrokenAI())
    db.refresh(p)
    assert p.status == Status.ACCEPTED and len(disk.files) == 1
    assert not media_path(p.id).exists()


def test_corrupt_file_rejected(db):
    setup_chat(db)
    p = intake_photo(db)
    mx = FakeMax()
    mx.content = b"not an image"
    process_photo(p.id, mx, FakeDisk(), GoodOCR(), GoodAI())
    db.refresh(p)
    assert p.status == Status.REJECTED and p.sha256 is None
    assert p.review_reason == ["CORRUPT_IMAGE"] and not media_path(p.id).exists()


def test_auth_roles_csrf_and_revocation(client, users, db):
    assert client.get("/api/photos").status_code == 401
    assert client.post("/api/auth/login", json={"login": "operator", "password": "wrong"}).status_code == 401
    assert (
        client.post(
            "/api/auth/login", json={"login": "operator", "password": "test-password-123"}
        ).status_code
        == 200
    )
    client.headers["X-CSRF-Token"] = client.cookies["csrf"]
    for path in ["/api/chats", "/api/users", "/api/audit", "/api/jobs"]:
        assert client.get(path).status_code == 403
    assert client.get("/api/photos").status_code == 200
    assert client.post("/api/chats", json={"max_chat_id": 1, "name": "X"}).status_code == 403
    assert client.get("/api/districts").status_code == 200
    client.headers.pop("X-CSRF-Token")
    assert client.post("/api/auth/logout").status_code == 403
    client.headers["X-CSRF-Token"] = client.cookies["csrf"]
    assert client.post("/api/auth/logout").status_code == 200
    assert client.get("/api/auth/me").status_code == 401


def test_review_version_audit_sync_and_dataset(admin, db, monkeypatch):
    setup_chat(db)
    setup_district(db)
    p = intake_photo(db)
    disk = FakeDisk()
    process_photo(p.id, FakeMax(), disk, GoodOCR(), MultipleAI())
    db.refresh(p)
    old_path = p.yandex_disk_path
    version = p.version
    initial_ai = p.ai_result
    result = admin.patch(
        f"/api/photos/{p.id}", json={"version": version, "photo_date": "2026-10-01", "work_type": "SWEEPING"}
    )
    assert result.status_code == 200, result.text
    assert result.json()["status"] == "ACCEPTED"
    assert (
        admin.patch(f"/api/photos/{p.id}", json={"version": version, "work_type": "BIN_EMPTYING"}).status_code
        == 409
    )
    monkeypatch.setattr("app.processing.YandexDisk", lambda: disk)
    sync_photo(p.id)
    db.refresh(p)
    assert p.work_type_source == "OPERATOR" and p.reviewed_by
    assert p.ai_result == initial_ai  # Operator ground truth does not overwrite AI prediction.
    assert p.yandex_disk_path != old_path and "/2026-10-01/" in p.yandex_disk_path
    assert old_path not in disk.files
    assert db.scalar(select(func.count()).select_from(Audit).where(Audit.photo_id == p.id)) == 1


def test_shared_filters_reports_list_and_excel(admin, db):
    setup_chat(db)
    setup_district(db)
    p1 = intake_photo(db)
    p2 = intake_photo(db, mid="second")
    process_photo(p1.id, FakeMax("red"), FakeDisk(), GoodOCR(), GoodAI())
    process_photo(p2.id, FakeMax("blue"), FakeDisk(), BrokenOCR(), GoodAI())
    for q, expected in [
        ("date_from=2026-09-30&date_to=2026-09-30&district_id=1", 1),
        ("status=NEEDS_REVIEW&district_id=0", 1),
        ("work_type=BIN_EMPTYING", 2),
        ("status=DUPLICATE", 0),
    ]:
        report = admin.get("/api/reports?" + q)
        photos = admin.get("/api/photos?" + q)
        assert report.status_code == 200, report.text
        assert photos.status_code == 200, photos.text
        assert report.json()["total"] == photos.json()["total"] == expected
        exported = admin.get("/api/export.xlsx?" + q)
        assert exported.status_code == 200, exported.text
        wb = load_workbook(BytesIO(exported.content), read_only=True)
        assert len(list(wb["Все фото"].rows)) == expected + 1
        assert len(wb.sheetnames) == 5
    assert admin.get("/api/photos?date_from=2026-10-02&date_to=2026-10-01").status_code == 422


def test_excel_formula_injection(admin, db):
    setup_chat(db)
    p = intake_photo(db)
    p.chat_name = '=HYPERLINK("https://example.com")'
    db.commit()
    xlsx = admin.get("/api/export.xlsx")
    assert xlsx.status_code == 200, xlsx.text
    wb = load_workbook(BytesIO(xlsx.content))
    assert wb["Все фото"]["E2"].data_type == "s"
    assert wb["Все фото"]["E2"].value.startswith("'=")


def test_polygon_overlap_confirmation_and_regeo(admin, db):
    setup_chat(db)
    d = setup_district(db)
    p = intake_photo(db)
    process_photo(p.id, FakeMax(), FakeDisk(), GoodOCR(), GoodAI())
    body = {
        "name": "7 мкр",
        "geometry": {
            "type": "Polygon",
            "coordinates": [[[72.5, 61], [73.5, 61], [73.5, 62], [72.5, 62], [72.5, 61]]],
        },
    }
    rejected = admin.post("/api/districts", json=body)
    assert rejected.status_code == 409, rejected.text
    assert rejected.json()["detail"]["overlaps"][0]["id"] == d.id
    allowed = admin.post("/api/districts", json={**body, "confirm_overlap": True})
    assert allowed.status_code == 201, allowed.text
    regeo_all()
    db.refresh(p)
    assert p.status == Status.NEEDS_REVIEW and p.district_id is None
    assert "DISTRICT_OVERLAP" in p.review_reason
    assert db.scalar(select(func.count()).select_from(Job).where(Job.kind == "SYNC")) >= 1


def test_last_admin_and_password_hash(admin, db):
    from app.models import User

    record = db.scalar(select(User).where(User.login == "admin"))
    assert record.password_hash.startswith("$argon2")
    assert admin.patch("/api/users/" + str(record.id), json={"active": False}).status_code == 409
    created = admin.post(
        "/api/users", json={"login": "new_user", "name": "Новый оператор", "password": "new-password-123"}
    )
    assert created.status_code == 201, created.text
    assert "password_hash" not in created.json()


def test_rate_limit_and_origin(client, users):
    body = {"login": "operator", "password": "incorrect"}
    assert (
        client.post("/api/auth/login", json=body, headers={"Origin": "https://evil.example"}).status_code
        == 403
    )
    for _ in range(8):
        assert client.post("/api/auth/login", json=body).status_code == 401
    assert client.post("/api/auth/login", json=body).status_code == 429


def test_worker_reclaims_running_job_and_retries(db, monkeypatch):
    from app.worker import run_once

    setup_chat(db)
    p = intake_photo(db)
    job = db.scalar(select(Job))
    job.status = "RUNNING"
    db.commit()
    calls = []
    monkeypatch.setattr("app.worker.process_batch", lambda oid: calls.append(oid))
    assert run_once()
    db.refresh(job)
    assert job.status == "DONE" and calls == [p.inbox_id]
    job.status = "PENDING"
    job.next_run_at = now() - timedelta(seconds=2)
    db.commit()

    def fail(_):
        raise RuntimeError("not persisted")

    monkeypatch.setattr("app.worker.process_batch", fail)
    assert run_once()
    db.refresh(job)
    assert job.status == "PENDING" and job.next_run_at > now() and job.last_error == "RuntimeError"


def test_album_one_aggregate_outbox(db, monkeypatch):
    from app.processing import process_batch

    setup_chat(db)
    setup_district(db)
    p = intake_photo(
        db,
        attachments=[
            {"type": "image", "payload": {"url": "https://cdn.max.ru/a"}},
            {"type": "file", "filename": "file.pdf"},
        ],
    )
    monkeypatch.setattr("app.processing.MaxClient", FakeMax)
    monkeypatch.setattr("app.processing.YandexDisk", FakeDisk)
    monkeypatch.setattr("app.processing.TesseractOCR", GoodOCR)
    monkeypatch.setattr("app.processing.LocalClassifier", GoodAI)
    process_batch(p.inbox_id)
    process_batch(p.inbox_id)
    inbox = db.get(Inbox, p.inbox_id)
    assert inbox.reply_text == "✅ Обработано фотографий: 2\nПринято: 1\nНе принято: 1"
    assert db.scalar(select(func.count()).select_from(Job).where(Job.kind == "REPLY")) == 1
