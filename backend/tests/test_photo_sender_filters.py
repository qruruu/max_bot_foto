"""The sender filter must share the photo scope without exposing other chats' participants."""

from io import BytesIO

import pytest
from openpyxl import load_workbook
from test_chat_access import AccessSession, actor
from test_chat_access import access_api as access_api
from test_chat_senders import LARGE_ID
from test_chat_senders import senders_api as senders_api

from app.models import Photo, User


def test_sender_options_use_shared_names_exact_ids_and_current_chat_access(senders_api):
    client, engine = senders_api
    with AccessSession(engine) as db:
        db.get(User, 2).permissions = ["photos_view"]
        db.commit()
    response = client.get('/api/photos/filter-counts?work_only=true')
    assert response.status_code == 200
    facet = response.json()['sender_id']
    assert facet == {
        'total': 2, 'values': {str(LARGE_ID): 1, '321': 1},
        'options': [{'value': str(LARGE_ID), 'name': 'Иван · бригада № 2'},
                    {'value': '321', 'name': None}],
    }
    assert 'Ещё не присылал фото' not in response.text
    assert 'Отправитель из другого чата' not in response.text
    assert '111' not in facet['values']  # Spam is outside the normal selection.


def test_sender_selection_keeps_other_filters_and_does_not_disclose_a_foreign_sender(senders_api):
    client, _ = senders_api
    base = client.get('/api/photos/filter-counts').json()
    selected = client.get('/api/photos/filter-counts', params={'sender_id': str(LARGE_ID)}).json()
    assert selected['sender_id'] == base['sender_id']
    assert selected['work_type']['total'] == 1
    assert client.get('/api/photos?sender_id=777').json()['total'] == 0
    assert client.get('/api/reports?sender_id=777').json()['total'] == 0
    foreign_chat = client.get('/api/photos/filter-counts?chat_id=200').json()['sender_id']
    assert foreign_chat == {'total': 0, 'values': {}, 'options': []}


def test_sender_applies_to_list_map_reports_export_next_and_return_position(senders_api):
    client, _ = senders_api
    actor(client, 1)
    params = {'sender_id': str(LARGE_ID)}
    photos = client.get('/api/photos', params=params).json()
    assert photos['total'] == 2
    assert {photo['id'] for photo in photos['items']} == {10, 20}
    assert {photo['sender_id'] for photo in photos['items']} == {str(LARGE_ID)}
    mapped = client.get('/api/photos/map', params=params).json()
    assert mapped['summary']['total'] == 2 and {p['id'] for p in mapped['items']} == {10, 20}
    assert client.get('/api/reports', params=params).json()['total'] == 2
    workbook = load_workbook(BytesIO(client.get('/api/export.xlsx', params=params).content))
    assert {row[0] for row in workbook['Все фото'].iter_rows(min_row=2, values_only=True)} == {10, 20}
    assert client.get('/api/photos/20/next', params=params).json()['id'] == 10
    position = client.get('/api/photos', params={**params, 'page_size': 1, 'anchor_id': 10}).json()
    assert position['page'] == 2 and position['focus_id'] == 10
    actor(client, 2)
    assert client.get('/api/photos', params=params).json()['total'] == 1


def test_missing_sender_is_distinct_from_an_unsigned_id_and_spam_stays_protected(senders_api):
    client, _ = senders_api
    actor(client, 1)
    assert client.get('/api/photos?sender_id=0').json()['total'] == 0
    missing = client.get('/api/photos?sender_id=0&status=SPAM').json()
    assert {photo['id'] for photo in missing['items']} == {40}
    assert client.get('/api/photos?sender_id=321').json()['items'][0]['id'] == 50


def test_sender_options_follow_current_labels_and_exclude_removed_previews(senders_api):
    client, engine = senders_api
    assert client.put(f'/api/chats/1/senders/{LARGE_ID}', json={'name': 'Новое имя'}).status_code == 200
    with AccessSession(engine) as db:
        photo = db.get(Photo, 50)
        photo.preview_removed_at = photo.received_at
        db.commit()
    facet = client.get('/api/photos/filter-counts').json()['sender_id']
    assert facet['options'] == [{'value': str(LARGE_ID), 'name': 'Новое имя'}]
    assert facet['total'] == 1


@pytest.mark.parametrize('value', ['-1', str(2**63), '1.5', 'invalid'])
def test_invalid_sender_filter_is_rejected(senders_api, value):
    client, _ = senders_api
    for path in ['/api/photos', '/api/photos/filter-counts', '/api/reports']:
        assert client.get(path, params={'sender_id': value}).status_code == 422
