"""Protect the existing server's catalog and model while keeping new installations empty."""

import json
import runpy
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from alembic.migration import MigrationContext
from alembic.operations import Operations
from openpyxl import load_workbook
from sqlalchemy import create_engine, delete, select, text
from sqlalchemy.orm import Session
from test_chat_access import AccessSession, actor
from test_chat_access import access_api as access_api

from app import ml
from app.catalog import WORK_TYPE_DISPLAY_ORDER, WORK_TYPES
from app.domain import storage_path
from app.exchange import PhotoRecord
from app.models import Audit, Chat, Photo, WorkType
from app.work_types import work_catalog

MIGRATION = Path(__file__).parents[1] / 'alembic/versions/0022_editable_work_types.py'


@pytest.mark.parametrize('existing_table', [None, 'users', 'chats', 'photos', 'model_versions'])
def test_migration_seeds_only_existing_installations_without_changing_data(existing_table, tmp_path):
    migration = runpy.run_path(str(MIGRATION))
    engine = create_engine('sqlite://')
    weights = tmp_path / 'weights.pt'
    weights.write_bytes(b'existing model weights must stay byte-identical')
    payload = json.dumps({'code': 'BIN_EMPTYING', 'labels': {'1': {'label': 'BIN_EMPTYING'}},
                          'status': 'ACTIVE', 'path': str(weights)})
    with engine.begin() as conn:
        for table in ['users', 'chats', 'photos', 'model_versions']:
            conn.execute(text(f'CREATE TABLE {table} (id INTEGER PRIMARY KEY, payload TEXT)'))
        if existing_table:
            conn.execute(text(f'INSERT INTO {existing_table} VALUES (1, :payload)'), {'payload': payload})
        before = {t: conn.execute(text(f'SELECT * FROM {t}')).all()
                  for t in ['users', 'chats', 'photos', 'model_versions']}
        with Operations.context(MigrationContext.configure(conn)):
            migration['upgrade']()
        after = {t: conn.execute(text(f'SELECT * FROM {t}')).all() for t in before}
        assert before == after
        rows = conn.execute(text('SELECT code,name,filename_label,model_index FROM work_types ORDER BY model_index')).all()
        assert rows == ([(code, names[0], names[1], i) for i, (code, names) in enumerate(WORK_TYPES.items())]
                        if existing_table else [])
        display = conn.execute(text('SELECT code FROM work_types ORDER BY sort_order')).scalars().all()
        assert display == (list(WORK_TYPE_DISPLAY_ORDER) if existing_table else [])
        # Simulate later reads/restarts: user edits remain durable and no boot seeding exists.
        if existing_table:
            conn.execute(text("UPDATE work_types SET name='Пользовательское название' WHERE code='BIN_EMPTYING'"))
    with Session(engine) as db:
        assert len(work_catalog(db)) == (16 if existing_table else 0)
        if existing_table:
            assert work_catalog(db)['BIN_EMPTYING'][0] == 'Пользовательское название'
    assert weights.read_bytes() == b'existing model weights must stay byte-identical'
    with pytest.raises(RuntimeError, match='preserved'):
        migration['downgrade']()
    engine.dispose()


def test_empty_catalog_stays_empty_and_unknown_legacy_assignments_are_rejected(access_api):
    client, engine = access_api
    actor(client, 1)
    with AccessSession(engine) as db:
        db.execute(delete(WorkType))
        db.commit()
    for _ in range(2):
        assert client.get('/api/meta').json()['work_types'] == []
        assert client.get('/api/work-types').json() == []
    assert client.get('/api/photos?work_type=BIN_EMPTYING').status_code == 422
    assert client.patch('/api/photos/10', json={'version': 1, 'work_type': 'BIN_EMPTYING'}).status_code == 422
    with AccessSession(engine) as db:
        assert not list(db.scalars(select(WorkType)))
    created = client.post('/api/work-types', json={'name': 'Первая работа', 'sort_order': 1})
    assert created.status_code == 201
    with AccessSession(engine) as db:
        row = db.get(WorkType, created.json()['id'])
        assert row.model_index == 0
        assert row.code.startswith('CUSTOM_')


def test_add_rename_and_reorder_keep_codes_photos_and_model_indices(access_api):
    client, engine = access_api
    actor(client, 1)
    original = client.get('/api/work-types').json()
    urns = next(row for row in original if row['code'] == 'BIN_EMPTYING')
    with AccessSession(engine) as db:
        legacy_indices = {w.code: w.model_index for w in db.scalars(select(WorkType))}
        original_path = storage_path(db.get(Photo, 10), None, work_catalog(db))
    response = client.put(f"/api/work-types/{urns['id']}", json={
        'name': 'Урны и корзины', 'sort_order': 50, 'version': urns['version']})
    assert response.status_code == 200 and response.json()['code'] == 'BIN_EMPTYING'
    assert response.json()['version'] == 2
    created = client.post('/api/work-types', json={'name': 'Новая работа', 'sort_order': 1})
    assert created.status_code == 201
    new = created.json()
    assert new['code'] not in WORK_TYPES
    assert client.get('/api/photos/10').json()['work_type_name'] == 'Урны и корзины'
    assert client.get('/api/photos/map').json()['items'][0]['work_type_name'] == 'Урны и корзины'
    assert client.get('/api/reports').json()['work_types'][0]['name'] == 'Урны и корзины'
    wb = load_workbook(BytesIO(client.get('/api/export.xlsx').content))
    assert wb['Все фото'].cell(2, 7).value == 'Урны и корзины'
    assert len(client.get('/api/meta').json()['work_types']) == 17
    with AccessSession(engine) as db:
        assert db.get(Photo, 10).work_type == 'BIN_EMPTYING'
        assert storage_path(db.get(Photo, 10), None, work_catalog(db)) == original_path
        assert {w.code: w.model_index for w in db.scalars(select(WorkType)) if w.code in WORK_TYPES} == legacy_indices
        assert db.get(WorkType, new['id']).model_index == 16
        assert db.scalar(select(Audit).where(Audit.action == 'WORK_TYPE_UPDATED')).old_value['name'] == 'УРНЫ'
    # The new category works in all the same workflows as the migrated ones.
    with AccessSession(engine) as db:
        db.get(Photo, 10).latitude = db.get(Photo, 10).longitude = None
        db.commit()
    saved = client.patch('/api/photos/10', json={'version': 1, 'work_type': new['code']})
    assert saved.status_code == 200 and saved.json()['work_type_name'] == 'Новая работа'
    for path in ['/api/photos', '/api/photos/map', '/api/photos/filter-counts', '/api/reports', '/api/export.xlsx']:
        assert client.get(path, params={'work_type': new['code']}).status_code == 200
    assert client.put('/api/chats/1', json={'max_chat_id': '100', 'name': 'Разрешённый',
        'cv_work_type_enabled': False, 'fixed_work_type_code': new['code']}).status_code == 200
    with AccessSession(engine) as db:
        photo = db.get(Photo, 10)
        assert 'Новая_работа' in storage_path(photo, None, work_catalog(db))
        assert db.get(Chat, 1).fixed_work_type_code == new['code']
        from app.exchange import record
        assert record(photo).source_work_type_code == new['code']
        assert PhotoRecord.model_validate(record(photo).model_dump()).source_work_type_code == new['code']


def test_catalog_writes_require_admin_validate_names_and_prevent_stale_or_code_changes(access_api):
    client, _ = access_api
    row = client.get('/api/work-types').json()[0]
    assert client.post('/api/work-types', json={'name': 'Нельзя', 'sort_order': 1}).status_code == 403
    assert client.put(f"/api/work-types/{row['id']}", json={**row, 'name': 'Нельзя'}).status_code == 403
    actor(client, 1)
    for name in ['', '   ', 'x' * 101, 'a\nb']:
        assert client.post('/api/work-types', json={'name': name, 'sort_order': 1}).status_code == 422
    assert client.post('/api/work-types', json={'name': 'УРНЫ', 'sort_order': 1}).status_code == 409
    assert client.post('/api/work-types', json={'name': 'Новая', 'sort_order': 0}).status_code == 422
    body = {'name': row['name'], 'sort_order': 2, 'version': row['version']}
    assert client.put(f"/api/work-types/{row['id']}", json={**body, 'code': 'REPLACED'}).status_code == 422
    assert client.put(f"/api/work-types/{row['id']}", json=body).status_code == 200
    assert client.put(f"/api/work-types/{row['id']}", json=body).status_code == 409
    assert client.delete(f"/api/work-types/{row['id']}").status_code == 405
    assert client.get('/api/photos?work_type=NONEXISTENT').status_code == 422
    assert client.post('/api/chats', json={'max_chat_id': '300', 'name': 'Новый',
        'cv_work_type_enabled': False, 'fixed_work_type_code': 'NONEXISTENT'}).status_code == 422


def test_legacy_and_new_model_heads_have_independent_immutable_code_maps(tmp_path):
    legacy = SimpleNamespace(metrics={'architecture': ml.ARCHITECTURE})
    assert ml.model_codes(legacy) == tuple(WORK_TYPES)
    codes = tuple(WORK_TYPES) + ('CUSTOM_NEW',)
    version = SimpleNamespace(metrics={'architecture': ml.ARCHITECTURE, 'class_codes': list(codes)})
    assert ml.model_codes(version) == codes
    torch, _, _ = ml.torch_modules()
    for model_codes in [ml.model_codes(legacy), codes, ('CUSTOM_FIRST', 'CUSTOM_SECOND')]:
        model = ml.network(None, model_codes)
        path = tmp_path / f'weights-{len(model_codes)}.pt'
        torch.save(model.state_dict(), path)
        digest = ml.digest_file(path)
        restored = ml.load_model(str(path), digest, model_codes)
        assert restored.fc.out_features == len(model_codes)
        assert torch.equal(model.fc.weight, restored.fc.weight)
        assert ml.digest_file(path) == digest
        row = [0.0] * len(model_codes)
        row[-1] = 1.0
        result = ml.interpret_probabilities([row], [model_codes[-1]], 0.8, model_codes)
        assert result.work_type == model_codes[-1]
    # Adding a catalog class doesn't reinterpret a single legacy output.
    assert ml.model_codes(legacy) == tuple(WORK_TYPES)
    ml.load_model.cache_clear()


def test_training_plan_and_incumbent_metrics_include_new_categories():
    codes = ('CUSTOM_FIRST', 'CUSTOM_SECOND')
    examples = [ml.Example(i, f'{i:064x}', codes[i % 2], Path('unused')) for i in range(120)]
    plan = ml.training_plan(examples, {c: (c, c) for c in codes})
    assert plan['ready'] and plan['eligible_classes'] == list(codes)
    assert [c['code'] for c in plan['classes']] == list(codes)
    metrics = ml.metrics_for([0, 1, 1], [0, 1, 1], list(codes), codes)
    assert metrics['macro_f1'] == 1


def test_existing_active_model_still_classifies_after_catalog_changes(access_api, tmp_path, monkeypatch):
    from PIL import Image
    from sqlalchemy import JSON, Column, MetaData, Table
    from sqlalchemy.dialects.postgresql import JSONB

    from app.models import ModelVersion

    client, engine = access_api
    actor(client, 1)
    metadata = MetaData()
    Table('model_versions', metadata, *[
        Column(c.name, JSON() if isinstance(c.type, JSONB) else c.type, primary_key=c.primary_key)
        for c in ModelVersion.__table__.columns
    ])
    metadata.create_all(engine)
    torch, _, _ = ml.torch_modules()
    model = ml.network()
    with torch.no_grad():
        model.fc.weight.zero_()
        model.fc.bias.fill_(-10)
        model.fc.bias[ml.CODES.index('BIN_EMPTYING')] = 10
    path = tmp_path / 'existing-16-head.pt'
    torch.save(model.state_dict(), path)
    checksum = ml.digest_file(path)
    image = tmp_path / 'photo.jpg'
    Image.new('RGB', (240, 200), 'green').save(image)
    labels = {'10': {'label': 'BIN_EMPTYING', 'sha256': 'a' * 64}}
    with AccessSession(engine) as db:
        db.add(ModelVersion(id=1, status='ACTIVE', dataset_fingerprint='b' * 64,
                            sample_count=100, labels_snapshot=labels, supported_classes=['BIN_EMPTYING'],
                            metrics={'architecture': ml.ARCHITECTURE}, model_path=str(path), model_sha256=checksum))
        db.commit()
    monkeypatch.setattr(ml, 'engine', lambda: engine)
    monkeypatch.setattr(ml, 'Session', AccessSession)
    assert ml.LocalClassifier().classify(image).classification.work_type == 'BIN_EMPTYING'
    assert client.post('/api/work-types', json={'name': 'Дополнительная работа', 'sort_order': 1}).status_code == 201
    row = next(w for w in client.get('/api/work-types').json() if w['code'] == 'BIN_EMPTYING')
    assert client.put(f"/api/work-types/{row['id']}", json={
        'name': 'Новое название урн', 'sort_order': 17, 'version': row['version']}).status_code == 200
    result = ml.LocalClassifier().classify(image)
    assert result.classification.work_type == 'BIN_EMPTYING'
    assert list(result.raw['probabilities']) == ml.CODES
    assert result.raw['model_id'] == 1
    with AccessSession(engine) as db:
        active = ml.active_model(db)
        assert active.status == 'ACTIVE' and active.labels_snapshot == labels
        assert active.model_sha256 == checksum and active.model_path == str(path)
        assert active.metrics == {'architecture': ml.ARCHITECTURE}
    assert ml.digest_file(path) == checksum
    ml.load_model.cache_clear()
