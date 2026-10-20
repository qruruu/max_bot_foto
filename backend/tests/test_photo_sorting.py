"""Normal operator sorting omits auxiliary archives while preserving explicit access."""
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text
from test_chat_access import AccessSession, actor
from test_chat_access import access_api as access_api

from app.models import Photo
from app.schemas import PhotoBrowseFilters, PhotoFilters, PhotoListFilters, PhotoMapFilters


@pytest.fixture
def sorting_api(access_api):
    client, engine = access_api
    actor(client,1)
    with AccessSession(engine) as db:
        db.get(Photo,20).status = 'DUPLICATE'
        db.get(Photo,60).status = 'FORWARDED'
        db.get(Photo,50).status = 'NEEDS_REVIEW'
        db.commit()
    return client, engine


def test_normal_list_and_next_skip_duplicates_forwarded_and_spam(sorting_api):
    client, _ = sorting_api
    result = client.get('/api/photos').json()
    assert result['total'] == 2 and [p['id'] for p in result['items']] == [50,10]
    assert client.get('/api/photos?page_size=1&page=2').json()['items'][0]['id'] == 10
    assert client.get('/api/photos/60/next').json() == {'id':50}
    assert client.get('/api/photos/50/next').json() == {'id':10}
    assert client.get('/api/photos/10/next').json() == {'id':None}
    assert client.get('/api/photos/10/next?status=NEEDS_REVIEW').json() == {'id':50}


@pytest.mark.parametrize('status,ids', [('DUPLICATE',{20}),('FORWARDED',{60}),('SPAM',{30,40})])
def test_explicit_status_still_opens_each_separate_archive(sorting_api,status,ids):
    client, _ = sorting_api
    result = client.get('/api/photos?work_only=true&status='+status).json()
    assert {p['id'] for p in result['items']} == ids
    mapped = client.get('/api/photos/map?work_only=true&status='+status).json()
    assert {p['id'] for p in mapped['items']} == ids
    latest = max(ids)
    assert client.get(f'/api/photos/{latest}/next?work_only=true&status='+status).json()['id'] == (min(ids) if len(ids)>1 else None)


@pytest.mark.parametrize('scope', ['', '?work_only=true', '?work_only=false'])
def test_map_notifications_and_export_match_the_sorting_list(sorting_api, scope):
    client, _ = sorting_api
    mapped = client.get('/api/photos/map').json()
    assert mapped['summary']['total'] == 2 and {p['id'] for p in mapped['items']} == {10,50}
    assert client.get('/api/photos/arrivals').json()['photo_ids'] == [10,50]
    response = client.get('/api/export.xlsx'+scope)
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert {row[0] for row in list(book['Все фото'].values)[1:]} == {10,50}
    assert sum(row[1] for row in list(book['По чатам'].values)[1:]) == 2


def test_filter_counts_match_normal_sorting_and_keep_archived_alternatives_visible(sorting_api):
    client, _ = sorting_api
    counts = client.get('/api/photos/filter-counts?work_only=true').json()
    assert all(counts[facet]['total'] == 2 for facet in counts)
    assert counts['status']['values']['DUPLICATE'] == 1
    assert counts['status']['values']['FORWARDED'] == 1
    assert counts['status']['values']['SPAM'] == 2
    assert counts['status']['values']['ACCEPTED_ALL'] == 2
    selected = client.get('/api/photos/filter-counts?work_only=true&status=DUPLICATE').json()
    assert selected['work_type']['total'] == 1 and selected['status']['total'] == 2
    assert selected['chat_id']['values'] == {'200':1}


@pytest.mark.parametrize('scope', ['', '?work_only=true', '?work_only=false'])
def test_overview_keeps_only_separate_archive_counters_and_no_records_are_removed(sorting_api, scope):
    client, engine = sorting_api
    report = client.get('/api/reports'+scope).json()
    assert report['total'] == 2 and report['accepted'] == 2
    assert report['counts']['NEEDS_REVIEW'] == 1
    assert report['counts']['DUPLICATE'] == report['counts']['FORWARDED'] == 1
    assert report['districts'] == [{'id':0, 'name':'Не определён', 'count':2}]
    assert report['chats'] == [{'id':'100', 'name':'Разрешённый', 'count':2}]
    assert report['work_types'] == [{'id':'BIN_EMPTYING', 'name':'УРНЫ', 'count':2}]
    assert report['matrix'] == [{'chat_id':'100', 'district_id':0, 'count':2}]
    assert client.get('/api/photos?work_only=false').json()['total'] == 4
    with AccessSession(engine) as db:
        assert set(db.scalars(select(Photo.id))) == {10,20,30,40,50,60}


def test_only_auxiliary_events_do_not_create_an_unassigned_district_or_work_groups(sorting_api):
    client, engine = sorting_api
    with AccessSession(engine) as db:
        for id in (20,60):
            db.get(Photo,id).sender_id = 455424204
        db.commit()
    filters = {'sender_id':'455424204', 'date_from':'2026-10-18',
               'date_to':'2026-10-18', 'date_basis':'photo'}
    report = client.get('/api/reports', params=filters).json()
    assert report['total'] == report['accepted'] == report['counts']['NEEDS_REVIEW'] == 0
    assert report['counts']['DUPLICATE'] == report['counts']['FORWARDED'] == 1
    assert all(report[key] == [] for key in ('districts', 'chats', 'work_types', 'matrix'))
    for status, id in [('DUPLICATE',20), ('FORWARDED',60)]:
        selected = client.get('/api/photos', params={**filters, 'status':status, 'work_only':'true'}).json()
        assert selected['total'] == 1 and [p['id'] for p in selected['items']] == [id]


def test_report_groups_and_drilldowns_count_only_work_photos_even_without_a_district(sorting_api):
    client, engine = sorting_api
    with AccessSession(engine) as db:
        db.execute(text("INSERT INTO districts (id, name, active) VALUES (1, '12 мкр', true)"))
        db.get(Photo,10).district_id = 1
        db.get(Photo,50).work_type = None
        for id in (10,20,50,60):
            db.get(Photo,id).sender_id = 455424204
        db.commit()
    filters = {'sender_id':'455424204', 'date_from':'2026-10-18', 'date_to':'2026-10-18'}
    report = client.get('/api/reports', params=filters).json()
    assert report['total'] == 2 and report['counts']['NEEDS_REVIEW'] == 1
    assert {d['id']:d['count'] for d in report['districts']} == {0:1, 1:1}
    assert {w['id']:w['count'] for w in report['work_types']} == {'BIN_EMPTYING':1, 'NONE':1}
    for district in report['districts']:
        selected = client.get('/api/photos', params={**filters, 'district_id':district['id'], 'work_only':'true'}).json()
        assert selected['total'] == district['count']
        assert all(p['id'] in {10,50} for p in selected['items'])
    for chat in report['chats']:
        selected = client.get('/api/photos', params={**filters, 'chat_id':chat['id'], 'work_only':'true'}).json()
        assert selected['total'] == chat['count']
    for work in report['work_types']:
        selected = client.get('/api/photos', params={**filters, 'work_type':work['id'], 'work_only':'true'}).json()
        assert selected['total'] == work['count']
    for cell in report['matrix']:
        selected = client.get('/api/photos', params={**filters, 'chat_id':cell['chat_id'],
                              'district_id':cell['district_id'], 'work_only':'true'}).json()
        assert selected['total'] == cell['count']


@pytest.mark.parametrize('status,id', [('DUPLICATE',20), ('FORWARDED',60), ('SPAM',None)])
def test_explicit_auxiliary_status_does_not_restore_work_breakdowns(sorting_api, status, id):
    client, _ = sorting_api
    report = client.get('/api/reports', params={'status':status, 'work_only':'false'}).json()
    assert report['total'] == report['accepted'] == 0
    assert report['counts'][status] == (2 if status == 'SPAM' else 1)
    assert all(report[key] == [] for key in ('districts', 'chats', 'work_types', 'matrix'))
    if id is not None:
        response = client.get('/api/export.xlsx', params={'status':status})
        assert response.status_code == 200
        book = load_workbook(BytesIO(response.content))
        assert [row[0] for row in list(book['Все фото'].values)[1:]] == [id]
        for title in ('По микрорайонам','По чатам','По видам работ','Чат × микрорайон'):
            assert len(list(book[title].values)) == 1


def test_sorting_scope_does_not_bypass_chat_or_spam_permissions(sorting_api):
    client, engine = sorting_api
    actor(client,2)
    assert client.get('/api/photos?status=DUPLICATE').json()['total'] == 0
    assert client.get('/api/photos?status=FORWARDED').json()['total'] == 0
    report = client.get('/api/reports?work_only=true').json()
    assert report['counts']['DUPLICATE'] == report['counts']['FORWARDED'] == 0
    assert report['total'] == 2
    from app.models import User
    with AccessSession(engine) as db:
        db.get(User,2).permissions = ['photos_view']
        db.commit()
    assert client.get('/api/photos?work_only=true&status=SPAM').status_code == 403
    assert client.get('/api/photos/map?work_only=true&status=SPAM').status_code == 403


def test_browsing_defaults_are_distinct_from_overview_and_report_filters():
    assert PhotoBrowseFilters().work_only
    assert PhotoListFilters().work_only
    assert PhotoMapFilters().work_only
    assert not PhotoFilters().work_only
    assert not PhotoListFilters(work_only=False).work_only
