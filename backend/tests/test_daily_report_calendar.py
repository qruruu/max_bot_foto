from datetime import date
from types import SimpleNamespace

import pytest
from pydantic import ValidationError

from app.daily_report_calendar import season_for, work_period
from app.daily_report_presets import shashko_config
from app.daily_report_schemas import ReportingConfig, ReportingDocument, ReportingLine, ReportingWork
from app.daily_reports import build_lines


def work(**values):
    return ReportingWork(id="work", name="Уборка", **values)


def test_four_windows_include_every_day_and_short_and_leap_months():
    w = work(recurrence="four_windows")
    for month, last in [(2, 28), (4, 30), (5, 31)]:
        for d in range(1, last + 1):
            first, end = work_period(date(2026, month, d), w, "", [])
            expected_first = 1 if d <= 7 else 8 if d <= 14 else 15 if d <= 21 else 22
            assert first.day == expected_first and first <= date(2026, month, d) <= end
            assert end.day == (expected_first + 6 if expected_first < 22 else last)
    assert work_period(date(2028, 2, 29), w, "", [])[1] == date(2028, 2, 29)


def test_calendar_seasons_fixed_dates_windows_deadlines_and_events():
    assert season_for(date(2026, 6, 1), ReportingConfig()) == ""
    config = ReportingConfig(summer_start="05-01", summer_end="10-15")
    assert season_for(date(2026, 10, 15), config) == "summer"
    assert season_for(date(2026, 10, 16), config) == "winter"
    assert work_period(date(2026, 5, 8), work(recurrence="dates", dates=["05-08"]), "", [])
    assert work_period(date(2026, 5, 9), work(recurrence="dates", dates=["05-08"]), "", []) is None
    wash = work(recurrence="windows", windows=[{"start": "06-01", "end": "06-15"}])
    assert work_period(date(2026, 6, 15), wash, "", []) == (date(2026, 6, 1), date(2026, 6, 15))
    assert work_period(date(2026, 6, 16), wash, "", []) is None
    mow = work(recurrence="deadlines", dates=["06-12", "07-31", "08-30"], season="summer")
    assert work_period(date(2026, 7, 12), mow, "summer", []) == (date(2026, 6, 13), date(2026, 7, 31))
    assert work_period(date(2026, 7, 12), mow, "winter", []) is None
    assert work_period(date(2026, 10, 27), work(recurrence="event"), "winter", []) is None
    assert work_period(date(2026, 10, 27), work(recurrence="event"), "winter", ["work"])


def test_only_confirmed_quantities_reduce_a_cycle_without_erasing_unknown_or_zero():
    config = shashko_config()
    w = next(w for w in config.works if w.id == "playground-monthly")
    w.assignments[0].quantity = 100
    config.works = [w]
    past = [(date(2026, 10, 10), ReportingDocument(fact=[ReportingLine(work_id=w.id, area_id="area58", decision="include", quantity=40)]))]
    lines = build_lines(date(2026, 10, 11), config, "", [], past)
    assert lines[0].quantity == 60 and lines[1].quantity is None
    past[0][1].fact[0].cycle_complete = True
    assert [r.area_id for r in build_lines(date(2026, 10, 11), config, "", [], past)] == ["area60"]
    w.assignments[1].quantity = 0
    assert build_lines(date(2026, 10, 11), config, "", [], past)[0].quantity == 0
    assert build_lines(date(2026, 11, 1), config, "", [], past)[0].quantity == 100


def test_shashko_source_has_no_demo_norms_and_keeps_contract_territories():
    config = shashko_config()
    assert len(config.works) == 15
    assert config.profile.short_name == "ИП Шашко А. Д."
    assert config.summer_start == config.summer_end == ""
    assert "12, 16" in config.areas[0].name
    assert "12" not in config.areas[1].name
    assert all(a.quantity is None and not a.resources and a.annual_limit is None for w in config.works for a in w.assignments)
    assert next(w for w in config.works if w.id == "playground-snow").unit == "уточнить"


@pytest.mark.parametrize("change", [
    {"quantity": float("nan")}, {"quantity": -1}, {"quantity": float("inf")},
    {"quantity": 0.0001},
])
def test_invalid_quantities_cannot_enter_a_report(change):
    with pytest.raises(ValidationError):
        ReportingLine(work_id="w", area_id="a", **change)


def test_report_export_retains_more_than_twenty_rows_and_escapes_untrusted_values():
    from io import BytesIO

    from openpyxl import load_workbook

    from app.daily_report_export import export_html, export_xlsx
    config = shashko_config()
    config.profile.company_name = "<script>alert(1)</script>"
    config.profile.recipient = ["=HYPERLINK(\"https://invalid.test\")"]
    template = config.works[0]
    config.works = [template.model_copy(update={"id": f"w{i}", "name": f"Работа {i}"}) for i in range(30)]
    document = ReportingDocument(fact=[ReportingLine(work_id=w.id, area_id="area58", decision="include", quantity=i + 1) for i, w in enumerate(config.works)])
    row = SimpleNamespace(config=config.model_dump(), approved_config=config.model_dump(), draft=document.model_dump(), approved=document.model_dump(),
                          day=date(2026, 10, 18), cohort="production", version=1, approved_version=1, letter_numbers={"58": "МК-58-005"})
    wb = load_workbook(BytesIO(export_xlsx(row)), data_only=False)
    letter = wb.worksheets[0]
    values = [c.value for row in letter for c in row]
    assert all(f"Работа {i}" in values for i in range(30))
    assert not any(c.data_type == "f" for ws in wb for row in ws for c in row)
    assert letter.page_setup.fitToHeight == 0
    html = export_html(row)
    assert "<script>alert(1)</script>" not in html and "&lt;script&gt;" in html
    assert "Работа 29" in html and "МК-58-005" in html
