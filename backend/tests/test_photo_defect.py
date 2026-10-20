import runpy
from datetime import UTC, date, datetime
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from alembic.migration import MigrationContext
from alembic.operations import Operations
from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import create_engine, text

from app.api import edit_photo
from app.catalog import WORK_TYPES, work_type_label
from app.domain import storage_path
from app.models import Audit, District, Photo
from app.reports import export_excel, photo_dict
from app.schemas import PhotoEdit, PhotoFilters


@pytest.fixture
def defect_photo(monkeypatch):
    monkeypatch.setattr("app.api.active_cohort", lambda db: "production")
    monkeypatch.setattr("app.api.work_catalog", lambda db: WORK_TYPES)
    monkeypatch.setattr("app.api.require_work_type", lambda db, code: None)
    monkeypatch.setattr("app.reports.work_catalog", lambda db: WORK_TYPES)
    photo = Photo(id=7, version=1, district_id=12, district_source="OPERATOR",
                  district_name_snapshot="12 мкр", photo_date=date(2026, 10, 18),
                  received_at=datetime(2026, 10, 18, tzinfo=UTC),
                  work_type="BIN_EMPTYING", work_type_source="AI", is_spam=False, is_defect=False,
                  status="ACCEPTED", analyzed_at=datetime(2026, 10, 18, tzinfo=UTC),
                  preview_removed_at=None, extension=".jpg", cohort="production", review_reason=[])
    district = District(id=12, name="12 мкр", active=True)
    db = Mock()
    db.get.side_effect = lambda model, id: photo if model is Photo else district
    db.scalar.return_value = None
    return photo, db


def save(photo, db, **changes):
    return edit_photo(photo.id, PhotoEdit(version=photo.version, **changes), db,
                      SimpleNamespace(id=1, role="ADMIN_OPERATOR"))


def test_inline_defect_edit_preserves_category_source_and_queues_sync(defect_photo):
    photo, db = defect_photo
    result = save(photo, db, is_defect=True)
    assert result["is_defect"] and result["work_type_name"] == "УРНЫ · НЕИСПРАВНОСТЬ"
    assert photo.work_type == "BIN_EMPTYING" and photo.work_type_source == "AI"
    assert photo.status == "ACCEPTED" and photo.version == 2 and photo.reviewed_by == 1
    assert photo.storage_status == "PENDING"
    entry = next(call.args[0] for call in db.add.call_args_list if isinstance(call.args[0], Audit))
    assert entry.old_value["is_defect"] is False and entry.new_value["is_defect"] is True
    db.commit.assert_called_once()


def test_marker_survives_category_edits_and_can_be_cleared(defect_photo):
    photo, db = defect_photo
    save(photo, db, is_defect=True)
    result = save(photo, db, work_type="GREEN_CLEANING")
    assert photo.is_defect and result["work_type_name"] == "ГАЗОН · НЕИСПРАВНОСТЬ"
    result = save(photo, db, is_defect=False)
    assert not photo.is_defect and result["work_type_name"] == "ГАЗОН"


@pytest.mark.parametrize("code", WORK_TYPES)
def test_defect_suffix_works_for_every_category_and_disk_filename(defect_photo, code):
    photo, _ = defect_photo
    photo.work_type, photo.is_defect = code, True
    assert work_type_label(code, True) == WORK_TYPES[code][0] + " · НЕИСПРАВНОСТЬ"
    assert photo_dict(photo, {}, {})["work_type_name"] == WORK_TYPES[code][0] + " · НЕИСПРАВНОСТЬ"
    assert storage_path(photo, "12 мкр").endswith(f"12_мкр_{WORK_TYPES[code][1]}_НЕИСПРАВНОСТЬ_7.jpg")
    photo.is_defect = False
    assert storage_path(photo, "12 мкр").endswith(f"12_мкр_{WORK_TYPES[code][1]}_7.jpg")


def test_spam_clears_marker_and_cannot_be_marked_until_restored(defect_photo):
    photo, db = defect_photo
    save(photo, db, is_defect=True)
    save(photo, db, is_spam=True)
    assert not photo.is_defect and photo.is_spam and photo.work_type is None
    with pytest.raises(HTTPException) as error:
        save(photo, db, is_defect=True)
    assert error.value.status_code == 422 and not photo.is_defect
    result = save(photo, db, work_type="BIN_EMPTYING", is_defect=True)
    assert photo.is_defect and not photo.is_spam and result["work_type_name"] == "УРНЫ · НЕИСПРАВНОСТЬ"


def test_stale_or_unprocessed_photo_cannot_change_defect_marker(defect_photo):
    photo, db = defect_photo
    with pytest.raises(HTTPException) as error:
        edit_photo(photo.id, PhotoEdit(version=2, is_defect=True), db, SimpleNamespace(id=1, role="ADMIN_OPERATOR"))
    assert error.value.status_code == 409 and not photo.is_defect
    photo.analyzed_at = None
    with pytest.raises(HTTPException) as error:
        save(photo, db, is_defect=True)
    assert error.value.status_code == 409 and not photo.is_defect
    db.commit.assert_not_called()


def test_omitting_the_marker_in_a_patch_does_not_reset_it():
    assert "is_defect" not in PhotoEdit(version=1).model_dump(exclude_unset=True)
    assert PhotoEdit(version=1, is_defect=False).model_dump(exclude_unset=True)["is_defect"] is False
    with pytest.raises(ValidationError):
        PhotoEdit(version=1, is_defect=None)


def test_excel_photo_label_includes_the_marker(defect_photo, monkeypatch):
    photo, db = defect_photo
    photo.is_defect = True
    db.execute.return_value.all.return_value = []
    db.scalars.return_value = [photo]
    monkeypatch.setattr("app.reports.summarize", lambda *args: dict(chats=[], districts=[], work_types=[], matrix=[]))
    book = load_workbook(BytesIO(export_excel(db, PhotoFilters())))
    assert book["Все фото"].cell(2, 7).value == "УРНЫ · НЕИСПРАВНОСТЬ"


def test_migration_defaults_existing_and_new_photos_to_unmarked():
    migration = runpy.run_path(str(Path(__file__).parents[1] / "alembic/versions/0019_photo_defect.py"))
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE photos (id INTEGER PRIMARY KEY)"))
        connection.execute(text("INSERT INTO photos (id) VALUES (1)"))
        with Operations.context(MigrationContext.configure(connection)):
            migration["upgrade"]()
        connection.execute(text("INSERT INTO photos (id) VALUES (2)"))
        assert connection.execute(text("SELECT is_defect FROM photos ORDER BY id")).scalars().all() == [0, 0]
    engine.dispose()
