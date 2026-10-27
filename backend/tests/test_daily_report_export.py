import base64
from copy import deepcopy
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from PIL import Image
from pydantic import ValidationError

from app.daily_report_export import document_title, export_html, export_xlsx
from app.daily_report_presets import shashko_config
from app.daily_report_schemas import ReportingProfile


def image_data(color):
    output = BytesIO()
    Image.new("RGB", (30, 10), color).save(output, "PNG")
    return "data:image/png;base64," + base64.b64encode(output.getvalue()).decode()


def report():
    config = shashko_config().model_dump()
    config["profile"]["letterhead_image"] = image_data("black")
    config["profile"]["signature_image"] = image_data("blue")
    document = {"fact": [{"work_id": "bins", "area_id": "area58", "decision": "info", "note": "Примечание оператора"}],
                "plan": [{"work_id": "bins", "area_id": "area58", "decision": "include", "quantity": 120}]}
    return SimpleNamespace(config=deepcopy(config), approved_config=deepcopy(config),
                           draft=deepcopy(document), approved=deepcopy(document), cohort="production",
                           day=date(2026, 10, 27), version=3, approved_version=2,
                           letter_numbers={"58": "МК-58-001", "60": "МК-60-001"})


def media(content):
    with ZipFile(BytesIO(content)) as archive:
        return [archive.read(name) for name in archive.namelist() if name.startswith("xl/media/")]


def test_final_signature_uses_approved_snapshot_and_is_absent_even_from_draft_xlsx_media():
    item = report()
    old_signature = item.approved_config["profile"]["signature_image"]
    item.config["profile"]["signature_image"] = image_data("red")
    signature_bytes = base64.b64decode(old_signature.split(",")[1])
    final = export_html(item)
    assert old_signature in final and item.config["profile"]["signature_image"] not in final
    assert "Подпись и печать организации" in final
    assert old_signature not in export_html(item, draft=True)
    assert signature_bytes in media(export_xlsx(item))
    assert signature_bytes not in media(export_xlsx(item, draft=True))
    assert base64.b64decode(item.config["profile"]["signature_image"].split(",")[1]) not in media(export_xlsx(item, draft=True))


def test_test_documents_never_contain_real_signature():
    item = report()
    item.cohort = "test"
    assert item.config["profile"]["signature_image"] not in export_html(item)
    assert len(media(export_xlsx(item))) == 2  # Letterhead only, one per contract.
    assert "ТЕСТОВЫЙ ОТЧЁТ" in export_html(item)


def test_letters_open_first_keep_numbers_numeric_and_do_not_repeat_internal_volume_status():
    item = report()
    workbook = load_workbook(BytesIO(export_xlsx(item)))
    assert workbook.active.title.startswith("Письмо 1")
    assert workbook.worksheets[-1].title == "Сведения"
    assert workbook.worksheets[-1].sheet_state == "hidden"
    for sheet in workbook.worksheets[:2]:
        assert sheet.sheet_view.showGridLines is False
        assert str(sheet.page_setup.paperSize) == sheet.PAPERSIZE_A4
        assert len(sheet._images) == 2
    cells = [cell.value for row in workbook.worksheets[0] for cell in row]
    assert 120 in cells and "—" in cells
    assert "Примечание оператора" in [str(value).replace("\n", " ") for value in cells]
    assert "Без указания объёма" not in cells
    html = export_html(item)
    assert "Без указания объёма" not in html and "Без объёма" not in html
    assert "Примечание оператора" in html and "План, объём" in html


def test_download_title_identifies_company_workday_and_draft_without_unsafe_characters():
    item = report()
    assert document_title(item) == "Отчёт ИП Шашко А. Д. за 26.10.2026"
    assert document_title(item, draft=True).startswith("Черновик — Отчёт ")
    item.cohort = "test"
    assert "Тестовый" in document_title(item, draft=True)
    item.config["profile"]["short_name"] = 'ООО "Тест/Путь"\r\n:файл?'
    title = document_title(item, draft=True)
    assert all(char not in title for char in '/\\\r\n:"?')
    assert title in export_html(item, draft=True)


@pytest.mark.parametrize("value", ["https://example.com/signature.png", "data:image/svg+xml;base64,PHN2Zz4=",
                                   "data:image/png;base64,bm90IGFuIGltYWdl", "data:image/png;base64,***"])
def test_document_images_cannot_load_remote_content_or_active_formats(value):
    with pytest.raises(ValidationError):
        ReportingProfile(signature_image=value)


def test_document_images_accept_original_png_without_reencoding_and_limit_dimensions():
    value = image_data("blue")
    assert ReportingProfile(signature_image=value).signature_image == value
    output = BytesIO()
    Image.new("RGB", (5001, 1)).save(output, "PNG")
    with pytest.raises(ValidationError):
        ReportingProfile(letterhead_image="data:image/png;base64," + base64.b64encode(output.getvalue()).decode())


def test_signature_occupies_one_indivisible_row_and_large_letters_repeat_column_headers():
    item = report()
    template = item.config["works"][0]
    item.config["works"] = [{**deepcopy(template), "id": f"work{i}", "name": f"Контрольная работа {i}"} for i in range(35)]
    item.approved_config = deepcopy(item.config)
    item.approved = {"fact": [{"work_id": w["id"], "area_id": "area58", "decision": "include", "quantity": 100 + i,
                              "resources": "3 человека", "note": "Контрольное примечание"} for i, w in enumerate(item.config["works"])],
                     "plan": item.approved["plan"]}
    item.approved["plan"][0]["work_id"] = "work0"
    workbook = load_workbook(BytesIO(export_xlsx(item)))
    sheet = workbook.worksheets[0]
    assert len(sheet.row_breaks.brk) >= 2
    cells = [cell.value for row in sheet for cell in row]
    normalized = [str(value).replace("\n", " ") for value in cells]
    assert all(f"Контрольная работа {i}" in normalized for i in range(35))
    assert normalized.count("Наименование работ, услуг") >= 3
    assert "Продолжение" in " ".join(normalized)
    signature = sheet._images[-1]
    signature_row = signature.anchor._from.row + 1
    assert sheet.row_dimensions[signature_row].height >= signature.anchor.ext.cy / 12700 - 1
    assert max(r.height or 15 for r in sheet.row_dimensions.values()) <= 409


def test_document_keeps_large_precise_quantities_and_unbroken_column_labels():
    item = report()
    item.approved["plan"][0]["quantity"] = 123456.789
    html = export_html(item)
    assert "123456,789 шт." in html and "1.23457e" not in html
    workbook = load_workbook(BytesIO(export_xlsx(item)))
    sheet = workbook.worksheets[0]
    cells = [cell for row in sheet for cell in row]
    assert any(cell.value == 123456.789 and cell.data_type == "n" for cell in cells)
    assert sum(cell.value == "Примечания" for cell in cells) == 2
